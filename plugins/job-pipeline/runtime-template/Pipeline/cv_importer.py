from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import multiprocessing
import os
import posixpath
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import threading
import uuid
import zipfile
import zlib
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlsplit
from xml.etree import ElementTree

import fitz
from docx import Document
from PIL import Image, UnidentifiedImageError

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from Pipeline.io_utils import (
    ConfinementError,
    atomic_write_bytes,
    atomic_write_json,
    confined_mkdir,
    confined_path,
    exclusive_project_lock,
    is_symlink_or_reparse,
    path_without_link_ancestors,
)
from Pipeline.onboarding import OnboardingError, _validate_cv, load_runtime
from Pipeline.policy_guard import PolicyError, _stable_single_link_bytes


class CVImportError(ValueError):
    """Field-safe failure while importing a private master CV."""


MAX_CV_BYTES = 15 * 1024 * 1024
MAX_CV_PAGES = 4
MAX_PAGE_POINTS = 2_000.0
MAX_PAGE_PIXELS = 20_000_000
MAX_PDF_XREFS = 20_000
MAX_PDF_DECODED_STREAM_BYTES = 64 * 1024 * 1024
MAX_PDF_STREAMS = 4_096
MAX_PDF_STREAM_BYTES = 8 * 1024 * 1024
MAX_PDF_STREAM_RATIO = 100
MAX_WORKER_IPC_BYTES = 8 * 1024 * 1024
MAX_WORKER_ARTIFACT_BYTES = 32 * 1024 * 1024
MAX_LAYOUT_BLOCKS = 10_000
MAX_LAYOUT_SPANS = 50_000
MAX_LAYOUT_TEXT_CHARS = 1_000_000
MAX_LAYOUT_IMAGES = 500
MAX_SECTION_ITEMS = 200
MAX_STYLE_ITEMS = 100
MAX_FONT_ITEMS = 100
MAX_XML_BYTES = 4 * 1024 * 1024
MAX_RENDER_SECONDS = 60
MAX_WORKER_MEMORY_BYTES = 384 * 1024 * 1024
MAX_MEDIA_DIMENSION = 8_000
MAX_MEDIA_PIXELS = 20_000_000
MAX_MEDIA_DECODED_BYTES = 64 * 1024 * 1024
MAX_DOCX_MEMBERS = 2_048
MAX_DOCX_MEMBER_BYTES = 4 * 1024 * 1024
MAX_DOCX_TOTAL_BYTES = 8 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}
ALLOWED_BLOCK_TYPES = {"text", "image", "paragraph", "table", "section", "header", "footer"}
EXECUTABLE_MEMBER_SUFFIXES = {
    ".bat",
    ".cmd",
    ".com",
    ".dll",
    ".exe",
    ".hta",
    ".jar",
    ".js",
    ".msi",
    ".ps1",
    ".scr",
    ".vbs",
}
DANGEROUS_FIELD_CODES = {"DDE", "DDEAUTO", "INCLUDETEXT", "INCLUDEPICTURE", "LINK"}
SHA256 = re.compile(r"[0-9a-f]{64}")
PDF_ACTIVE_KEYS = {"AA", "JS", "JavaScript", "Launch", "OpenAction", "RichMedia", "RichMediaContent", "RichMediaSettings"}
MEDIA_TYPES = {
    ".bmp": ("BMP", "image/bmp"),
    ".gif": ("GIF", "image/gif"),
    ".jpeg": ("JPEG", "image/jpeg"),
    ".jpg": ("JPEG", "image/jpeg"),
    ".png": ("PNG", "image/png"),
    ".tif": ("TIFF", "image/tiff"),
    ".tiff": ("TIFF", "image/tiff"),
    ".webp": ("WEBP", "image/webp"),
}
OFFICE_RELATIONSHIP_PREFIX = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
HYPERLINK_RELATIONSHIP = OFFICE_RELATIONSHIP_PREFIX + "hyperlink"
ALLOWED_CONTENT_TYPES = {
    "application/xml",
    "application/vnd.ms-word.commentsExtensible+xml".casefold(),
    "application/vnd.ms-word.commentsIds+xml".casefold(),
    "application/vnd.ms-word.people+xml".casefold(),
    "application/vnd.ms-word.stylesWithEffects+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.custom-properties+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.customXmlProperties+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.extended-properties+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.theme+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.glossary+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml".casefold(),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.webSettings+xml".casefold(),
    "application/vnd.openxmlformats-package.core-properties+xml".casefold(),
    "application/vnd.openxmlformats-package.relationships+xml".casefold(),
    *(content_type for _format, content_type in MEDIA_TYPES.values()),
}
IMAGE_CONTENT_TYPES = {content_type for _format, content_type in MEDIA_TYPES.values()}
RELATIONSHIP_TARGET_TYPES = {
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties": {"application/vnd.openxmlformats-package.core-properties+xml"},
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail": IMAGE_CONTENT_TYPES,
    OFFICE_RELATIONSHIP_PREFIX + "extended-properties": {"application/vnd.openxmlformats-officedocument.extended-properties+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "custom-properties": {"application/vnd.openxmlformats-officedocument.custom-properties+xml"},
    "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects": {"application/vnd.ms-word.styleswitheffects+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "comments": {"application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "customXml": {"application/xml"},
    OFFICE_RELATIONSHIP_PREFIX + "customXmlProps": {"application/vnd.openxmlformats-officedocument.customxmlproperties+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "endnotes": {"application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "fontTable": {"application/vnd.openxmlformats-officedocument.wordprocessingml.fonttable+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "footer": {"application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "footnotes": {"application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "glossaryDocument": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document.glossary+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "header": {"application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "image": IMAGE_CONTENT_TYPES,
    OFFICE_RELATIONSHIP_PREFIX + "numbering": {"application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "officeDocument": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "settings": {"application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "styles": {"application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "theme": {"application/vnd.openxmlformats-officedocument.theme+xml"},
    OFFICE_RELATIONSHIP_PREFIX + "webSettings": {"application/vnd.openxmlformats-officedocument.wordprocessingml.websettings+xml"},
}
ALLOWED_RELATIONSHIP_TYPES = {
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
    "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects",
    *{
        OFFICE_RELATIONSHIP_PREFIX + suffix
        for suffix in (
            "comments",
            "customXml",
            "customXmlProps",
            "endnotes",
            "fontTable",
            "footer",
            "footnotes",
            "glossaryDocument",
            "header",
            "hyperlink",
            "image",
            "numbering",
            "officeDocument",
            "settings",
            "styles",
            "theme",
            "webSettings",
        )
    },
}


@dataclass(frozen=True)
class SourceCV:
    path: Path
    kind: str
    sha256: str
    size_bytes: int
    _device: int = field(repr=False)
    _inode: int = field(repr=False)
    _mtime_ns: int = field(repr=False)
    _ctime_ns: int = field(repr=False)


@dataclass(frozen=True)
class PageGeometry:
    page_number: int
    width_points: float
    height_points: float
    margin_top_points: float | None = None
    margin_right_points: float | None = None
    margin_bottom_points: float | None = None
    margin_left_points: float | None = None


@dataclass(frozen=True)
class LayoutBlock:
    page_number: int
    block_type: str
    text: str
    bbox: tuple[float, float, float, float]
    font_names: tuple[str, ...]
    font_sizes: tuple[float, ...]
    colors: tuple[int, ...]
    style_names: tuple[str, ...] = ()


@dataclass(frozen=True)
class LayoutDescriptor:
    source_sha256: str
    source_kind: str
    page_count: int
    pages: tuple[PageGeometry, ...]
    blocks: tuple[LayoutBlock, ...]
    section_order: tuple[str, ...]
    image_count: int
    _attestation: str = field(default="", repr=False, compare=False)


@dataclass(frozen=True)
class _ImportAttestation:
    project_root: Path
    source_path: str
    source_identity: tuple[object, ...]
    descriptor_digest: str


_ATTESTATION_LOCK = threading.Lock()
_ATTESTATIONS: dict[str, _ImportAttestation] = {}


def _descriptor_payload(descriptor: LayoutDescriptor) -> dict[str, Any]:
    payload = asdict(descriptor)
    payload.pop("_attestation", None)
    return payload


def _descriptor_digest(descriptor: LayoutDescriptor) -> str:
    canonical = json.dumps(
        _descriptor_payload(descriptor), ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest()


def _metadata_identity(metadata: os.stat_result) -> tuple[int, int, int, int, int]:
    return (
        int(metadata.st_dev),
        int(metadata.st_ino),
        int(metadata.st_size),
        int(metadata.st_mtime_ns),
        int(metadata.st_ctime_ns),
    )


def _source_identity(source: SourceCV) -> tuple[object, ...]:
    return (
        source.path,
        source.kind,
        source.sha256,
        source.size_bytes,
        source._device,
        source._inode,
        source._mtime_ns,
        source._ctime_ns,
    )


def _absolute_checked_path(path: Path) -> Path:
    try:
        return path_without_link_ancestors(Path(path))
    except (ConfinementError, OSError, ValueError) as error:
        raise CVImportError("master CV source has a symlink or reparse ancestor") from error


def _read_stable_source(path: Path) -> tuple[Path, bytes, os.stat_result]:
    checked = _absolute_checked_path(path)
    try:
        metadata = checked.lstat()
    except OSError as error:
        raise CVImportError("master CV source must be a regular file") from error
    if is_symlink_or_reparse(checked):
        raise CVImportError("master CV source has a symlink or reparse ancestor")
    if not stat.S_ISREG(metadata.st_mode):
        raise CVImportError("master CV source must be a regular file")
    if metadata.st_nlink != 1:
        raise CVImportError("master CV source cannot be a hardlink")
    if not 0 < metadata.st_size <= MAX_CV_BYTES:
        raise CVImportError(f"master CV size must be 1-{MAX_CV_BYTES} bytes")
    try:
        payload = _stable_single_link_bytes(checked, MAX_CV_BYTES)
    except PolicyError as error:
        raise CVImportError("master CV source is unsafe or changed while reading") from error
    if not payload or len(payload) > MAX_CV_BYTES:
        raise CVImportError(f"master CV size must be 1-{MAX_CV_BYTES} bytes")
    try:
        after = checked.lstat()
    except OSError as error:
        raise CVImportError("master CV source changed while reading") from error
    if (
        is_symlink_or_reparse(checked)
        or not stat.S_ISREG(after.st_mode)
        or after.st_nlink != 1
        or _metadata_identity(after) != _metadata_identity(metadata)
    ):
        raise CVImportError("master CV source changed while reading")
    return checked, payload, after


def _pdf_raw_streams(payload: bytes):
    cursor = count = 0
    while True:
        marker = payload.find(b"stream", cursor)
        if marker < 0:
            return
        cursor = marker + 6
        if cursor >= len(payload) or payload[cursor : cursor + 1] not in {b"\r", b"\n"}:
            continue
        dictionary_end = payload.rfind(b">>", max(0, marker - 16_384), marker)
        dictionary_start = payload.rfind(b"<<", max(0, marker - 16_384), dictionary_end)
        end = payload.find(b"endstream", cursor)
        if dictionary_start < 0 or dictionary_end < dictionary_start or end < 0:
            raise CVImportError("master CV PDF stream structure is invalid")
        count += 1
        if count > MAX_PDF_STREAMS:
            raise CVImportError("master CV PDF stream count exceeds limits")
        data_start = cursor
        if payload[data_start : data_start + 2] == b"\r\n":
            data_start += 2
        else:
            data_start += 1
        data_end = end
        while data_end > data_start and payload[data_end - 1 : data_end] in {b"\r", b"\n"}:
            data_end -= 1
        data = payload[data_start:data_end]
        if len(data) > MAX_PDF_STREAM_BYTES:
            raise CVImportError("master CV PDF raw stream exceeds limits")
        yield payload[dictionary_start : dictionary_end + 2], data
        cursor = end + 9


def _preflight_pdf_raw(payload: bytes, *, check_compression: bool) -> None:
    if not isinstance(payload, bytes) or not 1 <= len(payload) <= MAX_CV_BYTES or not payload.startswith(b"%PDF-"):
        raise CVImportError("master CV PDF magic is invalid")
    if payload.count(b" obj") > MAX_PDF_XREFS:
        raise CVImportError("master CV PDF object count exceeds limits")
    allowed_filters = {b"ASCII85Decode", b"ASCIIHexDecode", b"CCITTFaxDecode", b"DCTDecode", b"FlateDecode", b"JPXDecode", b"RunLengthDecode"}
    for dictionary, stream in _pdf_raw_streams(payload):
        filter_match = re.search(rb"/Filter\s*(\[[^\]]{0,256}\]|/[A-Za-z0-9]+)", dictionary)
        filters = tuple(re.findall(rb"/([A-Za-z0-9]+)", filter_match.group(1))) if filter_match else ()
        if any(item not in allowed_filters for item in filters) or len(filters) > 1:
            raise CVImportError("master CV PDF stream filter is not allowed")
        if b"/Subtype" in dictionary and b"/Image" in dictionary:
            width_match = re.search(rb"/Width\s+(\d+)", dictionary)
            height_match = re.search(rb"/Height\s+(\d+)", dictionary)
            if width_match is None or height_match is None:
                raise CVImportError("master CV PDF image dictionary is invalid")
            width, height = int(width_match.group(1)), int(height_match.group(1))
            if not 1 <= width <= MAX_MEDIA_DIMENSION or not 1 <= height <= MAX_MEDIA_DIMENSION or width * height > MAX_PAGE_PIXELS:
                raise CVImportError("master CV PDF image dimensions exceed limits")
        if check_compression and filters == (b"FlateDecode",):
            limit = min(MAX_PDF_DECODED_STREAM_BYTES, max(1_024, len(stream) * MAX_PDF_STREAM_RATIO))
            try:
                decoder = zlib.decompressobj()
                decoded = decoder.decompress(stream, limit + 1)
                if len(decoded) > limit or decoder.unconsumed_tail:
                    raise CVImportError("master CV PDF stream compression ratio exceeds limits")
                decoded += decoder.flush(limit - len(decoded) + 1)
                if len(decoded) > limit:
                    raise CVImportError("master CV PDF stream compression ratio exceeds limits")
            except CVImportError:
                raise
            except zlib.error as error:
                raise CVImportError("master CV PDF compressed stream is invalid") from error


def _pdf_structure_unsafe(payload: bytes) -> tuple[int, bool, bool]:
    try:
        with fitz.open(stream=payload, filetype="pdf") as document:
            if document.needs_pass:
                return 0, True, False
            embedded = bool(document.embfile_names())
            active = False
            if document.xref_length() > MAX_PDF_XREFS:
                raise CVImportError("master CV PDF exceeds parser limits")
            decoded_total = 0
            for xref in range(1, document.xref_length()):
                keys = document.xref_get_keys(xref)
                if not isinstance(keys, (tuple, list)) or any(not isinstance(key, str) for key in keys):
                    raise CVImportError("master CV PDF structure is invalid")
                if PDF_ACTIVE_KEYS.intersection(keys):
                    active = True
                for key in keys:
                    value_type, value = document.xref_get_key(xref, key)
                    if value_type == "name" and value in {"/JavaScript", "/Launch", "/RichMedia"}:
                        active = True
                    if value_type == "name" and value in {"/EmbeddedFile", "/FileAttachment", "/Filespec"}:
                        embedded = True
                if document.xref_is_stream(xref):
                    stream = document.xref_stream(xref)
                    decoded_total += len(stream or b"")
                    if decoded_total > MAX_PDF_DECODED_STREAM_BYTES:
                        raise CVImportError("master CV PDF exceeds parser limits")
            for page in document:
                for link in page.get_links():
                    kind = link.get("kind")
                    if kind not in {fitz.LINK_GOTO, fitz.LINK_URI}:
                        active = True
                    elif kind == fitz.LINK_URI:
                        scheme = urlsplit(str(link.get("uri", ""))).scheme.casefold()
                        if scheme not in {"http", "https", "mailto"}:
                            active = True
            return document.page_count, False, embedded or active
    except CVImportError:
        raise
    except Exception as error:
        raise CVImportError("master CV PDF is invalid") from error


def _validate_pdf_payload_unsafe(payload: bytes) -> None:
    if not payload.startswith(b"%PDF-"):
        raise CVImportError("master CV PDF magic is invalid")
    if b"%%EOF" not in payload and any(marker in payload for marker in (b"/OpenAction", b"/JavaScript", b"/Launch")):
        raise CVImportError("master CV PDF contains active content")
    page_count, encrypted, unsafe = _pdf_structure_unsafe(payload)
    if encrypted:
        raise CVImportError("master CV PDF is encrypted")
    if unsafe:
        try:
            with fitz.open(stream=payload, filetype="pdf") as document:
                if document.embfile_names():
                    raise CVImportError("master CV PDF contains embedded content")
        except CVImportError:
            raise
        except Exception:
            pass
        raise CVImportError("master CV PDF contains active content")
    if not 1 <= page_count <= MAX_CV_PAGES:
        raise CVImportError(f"master CV PDF page count must be 1-{MAX_CV_PAGES}")
    try:
        _validate_cv(".pdf", payload)
    except OnboardingError as error:
        raise CVImportError("master CV PDF is invalid") from error


def _xml_root(payload: bytes, label: str) -> ElementTree.Element:
    if len(payload) > MAX_XML_BYTES:
        raise CVImportError("master CV DOCX archive limits exceeded")
    if payload.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff", b"\xff\xfe", b"\xfe\xff")) or b"\x00" in payload:
        raise CVImportError("master CV DOCX XML encoding is not allowed")
    try:
        text = payload.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise CVImportError("master CV DOCX XML encoding is not allowed") from error
    declaration = re.match(r"\A\s*<\?xml\b(?P<attributes>.*?)\?>", text, flags=re.IGNORECASE | re.DOTALL)
    if declaration is not None:
        encoding = re.search(
            r"\bencoding\s*=\s*(['\"])(?P<name>[^'\"]+)\1",
            declaration.group("attributes"),
            flags=re.IGNORECASE,
        )
        if encoding is not None:
            name = encoding.group("name").casefold().replace("_", "-")
            if name not in {"utf-8", "utf8", "us-ascii"} or (name == "us-ascii" and any(ord(character) > 127 for character in text)):
                raise CVImportError("master CV DOCX XML encoding is not allowed")
    if re.search(r"<\s*!\s*(?:doctype|entity)\b", text, flags=re.IGNORECASE):
        raise CVImportError("master CV DOCX contains DTD or entity content")
    try:
        return ElementTree.fromstring(text)
    except ElementTree.ParseError as error:
        raise CVImportError(f"master CV DOCX {label} is invalid") from error


def _field_code(value: str) -> str | None:
    normalized = " ".join(value.upper().split())
    first = normalized.split(" ", 1)[0] if normalized else ""
    return first if first in DANGEROUS_FIELD_CODES else None


def _relationship_base(name: str) -> str:
    pure = PurePosixPath(name)
    if pure == PurePosixPath("_rels/.rels"):
        return ""
    parts = pure.parts
    if len(parts) < 3 or parts[-2] != "_rels" or not parts[-1].endswith(".rels"):
        return ""
    source_name = parts[-1][:-5]
    return PurePosixPath(*parts[:-2], source_name).parent.as_posix()


def _validate_relationships(name: str, root: ElementTree.Element, part_types: dict[str, str]) -> None:
    base = _relationship_base(name)
    for relationship in root:
        relation_type = relationship.attrib.get("Type", "")
        target = relationship.attrib.get("Target", "")
        mode = relationship.attrib.get("TargetMode", "")
        if not isinstance(relation_type, str) or not isinstance(target, str) or not isinstance(mode, str):
            raise CVImportError("master CV DOCX relationship is invalid")
        if len(target) > 2048 or "\\" in target or "\x00" in target:
            raise CVImportError("master CV DOCX relationship is invalid")
        if mode.casefold() == "external":
            parsed = urlsplit(target)
            if (
                relation_type != HYPERLINK_RELATIONSHIP
                or parsed.scheme.casefold() not in {"http", "https"}
                or not parsed.netloc
                or parsed.username is not None
                or parsed.password is not None
            ):
                raise CVImportError("master CV DOCX contains unsafe external relationship")
        elif mode:
            raise CVImportError("master CV DOCX relationship mode is invalid")
        else:
            if relation_type not in ALLOWED_RELATIONSHIP_TYPES:
                raise CVImportError("master CV DOCX relationship type is not allowed")
            parsed = urlsplit(target)
            normalized = posixpath.normpath(posixpath.join(base, parsed.path))
            if parsed.scheme or parsed.netloc or target.startswith("/") or normalized == ".." or normalized.startswith("../"):
                raise CVImportError("master CV DOCX relationship escapes package")
            normalized = normalized.casefold()
            if normalized not in part_types:
                raise CVImportError("master CV DOCX relationship target is missing")
            allowed_target_types = RELATIONSHIP_TARGET_TYPES.get(relation_type)
            if allowed_target_types is None or part_types[normalized] not in allowed_target_types:
                raise CVImportError("master CV DOCX relationship target type is not allowed")


def _content_type_maps(root: ElementTree.Element) -> tuple[dict[str, str], dict[str, str]]:
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for element in root:
        local = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
        content_type = element.attrib.get("ContentType", "")
        if not isinstance(content_type, str) or not content_type:
            raise CVImportError("master CV DOCX content type is invalid")
        if "macroenabled" in content_type.casefold() or "vbaproject" in content_type.casefold():
            raise CVImportError("master CV DOCX contains macro content type")
        if content_type.casefold() not in ALLOWED_CONTENT_TYPES:
            raise CVImportError("master CV DOCX content type is not allowed")
        if local == "Default":
            extension = element.attrib.get("Extension", "")
            if not isinstance(extension, str) or not extension or extension.casefold() in defaults:
                raise CVImportError("master CV DOCX content type is invalid")
            defaults[extension.casefold()] = content_type.casefold()
        elif local == "Override":
            part = element.attrib.get("PartName", "")
            if not isinstance(part, str) or not part.startswith("/") or part.casefold() in overrides:
                raise CVImportError("master CV DOCX content type is invalid")
            overrides[part.casefold()] = content_type.casefold()
        else:
            raise CVImportError("master CV DOCX content type is invalid")
    return defaults, overrides


def _validate_media(name: str, payload: bytes, content_type: str) -> None:
    suffix = PurePosixPath(name).suffix.casefold()
    expected = MEDIA_TYPES.get(suffix)
    if expected is None:
        raise CVImportError("master CV DOCX media type is not allowed")
    expected_format, expected_content_type = expected
    if content_type != expected_content_type:
        raise CVImportError("master CV DOCX media content type mismatch")
    try:
        with Image.open(io.BytesIO(payload)) as image:
            width, height = image.size
            if image.format != expected_format:
                raise CVImportError("master CV DOCX media signature mismatch")
            if (
                type(width) is not int
                or type(height) is not int
                or not 1 <= width <= MAX_MEDIA_DIMENSION
                or not 1 <= height <= MAX_MEDIA_DIMENSION
                or width * height > MAX_MEDIA_PIXELS
            ):
                raise CVImportError("master CV DOCX media dimensions exceed limits")
            bands = image.getbands()
            if not isinstance(bands, tuple) or not 1 <= len(bands) <= 4:
                raise CVImportError("master CV DOCX media signature mismatch")
            if width * height * max(4, len(bands)) > MAX_MEDIA_DECODED_BYTES:
                raise CVImportError("master CV DOCX media decoded size exceeds limits")
            image.verify()
    except CVImportError:
        raise
    except Image.DecompressionBombError as error:
        raise CVImportError("master CV DOCX media dimensions exceed limits") from error
    except (OSError, ValueError, SyntaxError, UnidentifiedImageError) as error:
        raise CVImportError("master CV DOCX media signature is invalid") from error


def _docx_extension(name: str) -> str:
    leaf = PurePosixPath(name).name
    return leaf.rsplit(".", 1)[1].casefold() if "." in leaf else ""


def _safe_part_path(name: str, content_type: str) -> bool:
    lowered = name.casefold()
    if content_type in IMAGE_CONTENT_TYPES:
        return True
    if content_type == "application/vnd.openxmlformats-package.relationships+xml":
        return lowered == "_rels/.rels" or bool(re.fullmatch(r"(?:.+/)?_rels/[^/]+\.rels", lowered))
    rules = {
        "application/xml": r"customxml/item\d+\.xml",
        "application/vnd.ms-word.commentsextensible+xml": r"word/commentsextended\.xml",
        "application/vnd.ms-word.commentsids+xml": r"word/commentsids\.xml",
        "application/vnd.ms-word.people+xml": r"word/people\.xml",
        "application/vnd.ms-word.styleswitheffects+xml": r"word/styleswitheffects\.xml",
        "application/vnd.openxmlformats-officedocument.custom-properties+xml": r"docprops/custom\.xml",
        "application/vnd.openxmlformats-officedocument.customxmlproperties+xml": r"customxml/itemprops\d+\.xml",
        "application/vnd.openxmlformats-officedocument.extended-properties+xml": r"docprops/app\.xml",
        "application/vnd.openxmlformats-officedocument.theme+xml": r"word/(?:glossary/)?theme/theme\d+\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml": r"word/comments\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.glossary+xml": r"word/glossary/document\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml": r"word/document\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml": r"word/endnotes\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.fonttable+xml": r"word/(?:glossary/)?fonttable\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml": r"word/footer\d+\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml": r"word/footnotes\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml": r"word/header\d+\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml": r"word/(?:glossary/)?numbering\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml": r"word/(?:glossary/)?settings\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml": r"word/(?:glossary/)?styles\.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.websettings+xml": r"word/(?:glossary/)?websettings\.xml",
        "application/vnd.openxmlformats-package.core-properties+xml": r"docprops/core\.xml",
    }
    pattern = rules.get(content_type)
    return pattern is not None and re.fullmatch(pattern, lowered) is not None


def _docx_metadata_preflight(archive: zipfile.ZipFile) -> tuple[list[zipfile.ZipInfo], dict[str, zipfile.ZipInfo]]:
    infos = archive.infolist()
    if not 1 <= len(infos) <= MAX_DOCX_MEMBERS:
        raise CVImportError("master CV DOCX archive limits exceeded")
    by_folded_name: dict[str, zipfile.ZipInfo] = {}
    total_compressed = total_uncompressed = 0
    for item in infos:
        name = item.filename
        pure = PurePosixPath(name)
        canonical = posixpath.normpath(name.rstrip("/"))
        if (
            not name
            or "\\" in name
            or pure.is_absolute()
            or ".." in pure.parts
            or "." in pure.parts
            or re.match(r"^[a-z]:", name, re.I)
            or canonical != name.rstrip("/")
            or (item.is_dir() and not name.endswith("/"))
            or (not item.is_dir() and name.endswith("/"))
            or stat.S_IFMT(item.external_attr >> 16) == stat.S_IFLNK
        ):
            raise CVImportError("master CV DOCX contains traversal content")
        folded = name.casefold()
        if folded in by_folded_name:
            raise CVImportError("master CV DOCX contains duplicate members")
        by_folded_name[folded] = item
        if item.is_dir():
            continue
        if item.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
            raise CVImportError("master CV DOCX compression method is not allowed")
        if item.flag_bits & 1:
            raise CVImportError("master CV DOCX is encrypted")
        total_compressed += item.compress_size
        total_uncompressed += item.file_size
        if (
            item.file_size < 0
            or item.compress_size < 0
            or item.file_size > MAX_DOCX_MEMBER_BYTES
            or item.compress_size > MAX_DOCX_MEMBER_BYTES
            or total_uncompressed > MAX_DOCX_TOTAL_BYTES
            or total_compressed > MAX_DOCX_TOTAL_BYTES
            or (item.file_size and (item.compress_size == 0 or item.file_size / item.compress_size > 100))
        ):
            raise CVImportError("master CV DOCX archive limits exceeded")
        if folded in {"encryptedpackage", "encryptioninfo"} or folded.endswith("/encryptedpackage"):
            raise CVImportError("master CV DOCX is encrypted")
        if (
            folded.startswith(("word/embeddings/", "word/activex/", "customui/"))
            or folded in {"word/vbaproject.bin", "word/vbadata.xml"}
            or PurePosixPath(folded).suffix in EXECUTABLE_MEMBER_SUFFIXES
        ):
            raise CVImportError("master CV DOCX contains active or embedded content")
    if "[Content_Types].xml" not in {item.filename for item in infos}:
        raise CVImportError("master CV DOCX content type is missing")
    return infos, by_folded_name


def _validate_docx_payload(payload: bytes) -> None:
    if not payload.startswith(b"PK\x03\x04"):
        raise CVImportError("master CV DOCX magic is invalid")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos, by_folded_name = _docx_metadata_preflight(archive)
            content_info = next(item for item in infos if item.filename == "[Content_Types].xml")
            content_root = _xml_root(archive.read(content_info), "content types")
            defaults, overrides = _content_type_maps(content_root)
            part_types: dict[str, str] = {}
            for item in infos:
                if item.is_dir() or item.filename == "[Content_Types].xml":
                    continue
                lowered = item.filename.casefold()
                content_type = overrides.get("/" + lowered, defaults.get(_docx_extension(lowered), ""))
                if not content_type or content_type not in ALLOWED_CONTENT_TYPES or not _safe_part_path(lowered, content_type):
                    raise CVImportError("master CV DOCX content type is not allowed for part")
                part_types[lowered] = content_type
                if content_type in IMAGE_CONTENT_TYPES:
                    _validate_media(item.filename, archive.read(item), content_type)
            if any(part[1:].casefold() not in by_folded_name for part in overrides):
                raise CVImportError("master CV DOCX content type override target is missing")
            for item in infos:
                if item.is_dir() or item.filename == "[Content_Types].xml":
                    continue
                lowered = item.filename.casefold()
                if part_types[lowered] not in ALLOWED_CONTENT_TYPES - IMAGE_CONTENT_TYPES:
                    continue
                root = _xml_root(archive.read(item), "XML")
                for element in root.iter():
                    local = element.tag.rsplit("}", 1)[-1].casefold() if isinstance(element.tag, str) else ""
                    if local in {"altchunk", "object", "oleobject"}:
                        raise CVImportError("master CV DOCX contains active content")
                    if local in {"instrtext", "fldsimple", "fldchar"}:
                        raise CVImportError("master CV DOCX contains dangerous field instruction")
                if lowered.endswith(".rels"):
                    _validate_relationships(item.filename, root, part_types)
    except CVImportError:
        raise
    except (KeyError, OSError, RuntimeError, NotImplementedError, zipfile.BadZipFile) as error:
        raise CVImportError("master CV DOCX is invalid") from error
    try:
        _validate_cv(".docx", payload)
    except OnboardingError as error:
        raise CVImportError("master CV DOCX is invalid") from error


def validate_source_cv(path: Path) -> SourceCV:
    candidate = Path(path)
    kind = ALLOWED_EXTENSIONS.get(candidate.suffix.casefold())
    if kind is None:
        raise CVImportError("master CV must be PDF or DOCX")
    checked, payload, metadata = _read_stable_source(candidate)
    if kind == "pdf":
        _validate_pdf_payload(payload)
    else:
        _validate_docx_payload(payload)
    return SourceCV(
        path=checked,
        kind=kind,
        sha256=hashlib.sha256(payload).hexdigest(),
        size_bytes=len(payload),
        _device=int(metadata.st_dev),
        _inode=int(metadata.st_ino),
        _mtime_ns=int(metadata.st_mtime_ns),
        _ctime_ns=int(metadata.st_ctime_ns),
    )


def _safe_work_dir(work_dir: Path) -> Path:
    try:
        checked = path_without_link_ancestors(Path(work_dir))
        if checked.exists() and (not checked.is_dir() or is_symlink_or_reparse(checked)):
            raise CVImportError("CV import work directory is unsafe")
        checked.mkdir(parents=True, exist_ok=True)
        checked = path_without_link_ancestors(checked)
        if not checked.is_dir() or is_symlink_or_reparse(checked):
            raise CVImportError("CV import work directory is unsafe")
        return checked
    except CVImportError:
        raise
    except (ConfinementError, OSError, ValueError) as error:
        raise CVImportError("CV import work directory is unsafe") from error


def _current_payload(source: SourceCV) -> bytes:
    if (
        type(source) is not SourceCV
        or source.kind not in {"pdf", "docx"}
        or SHA256.fullmatch(source.sha256) is None
        or type(source.size_bytes) is not int
        or not 1 <= source.size_bytes <= MAX_CV_BYTES
        or any(type(value) is not int for value in (source._device, source._inode, source._mtime_ns, source._ctime_ns))
    ):
        raise CVImportError("source CV identity is invalid")
    try:
        checked, payload, metadata = _read_stable_source(source.path)
    except CVImportError as error:
        raise CVImportError("source CV is unsafe or changed after validation") from error
    if (
        checked != source.path
        or len(payload) != source.size_bytes
        or hashlib.sha256(payload).hexdigest() != source.sha256
        or _metadata_identity(metadata)
        != (source._device, source._inode, source.size_bytes, source._mtime_ns, source._ctime_ns)
    ):
        raise CVImportError("source CV changed after validation")
    if ALLOWED_EXTENSIONS.get(checked.suffix.casefold()) != source.kind:
        raise CVImportError("source CV changed after validation")
    if source.kind == "pdf":
        _validate_pdf_payload(payload)
    else:
        _validate_docx_payload(payload)
    return payload


def _finite_number(value: object) -> float:
    if type(value) not in {int, float} or not math.isfinite(float(value)):
        raise CVImportError("layout contains non-finite geometry")
    return float(value)


def _geometry(page_number: int, width: object, height: object) -> PageGeometry:
    width_value, height_value = _finite_number(width), _finite_number(height)
    if not 1 <= width_value <= MAX_PAGE_POINTS or not 1 <= height_value <= MAX_PAGE_POINTS:
        raise CVImportError("master CV page geometry exceeds limits")
    if math.ceil(width_value * 2) * math.ceil(height_value * 2) > MAX_PAGE_PIXELS:
        raise CVImportError("master CV page render exceeds pixel limits")
    return PageGeometry(page_number, width_value, height_value)


def _ordered_unique(values: list[Any]) -> tuple[Any, ...]:
    return tuple(dict.fromkeys(values))


def _heading(text: str, maximum_size: float) -> str | None:
    candidate = " ".join(text.split())
    if not candidate or len(candidate) > 120 or "\n" in text.strip("\n"):
        return None
    letters = [character for character in candidate if character.isalpha()]
    if maximum_size >= 11 or (letters and all(character.isupper() for character in letters)):
        return candidate
    return None


def _pdf_layout_unsafe(payload: bytes, source_sha256: str, source_kind: str) -> LayoutDescriptor:
    pages: list[PageGeometry] = []
    blocks: list[LayoutBlock] = []
    sections: list[str] = []
    image_count = span_count = text_chars = 0
    try:
        with fitz.open(stream=payload, filetype="pdf") as document:
            if not 1 <= document.page_count <= MAX_CV_PAGES:
                raise CVImportError(f"master CV PDF page count must be 1-{MAX_CV_PAGES}")
            for page_index, page in enumerate(document, start=1):
                geometry = _geometry(page_index, page.rect.width, page.rect.height)
                pages.append(geometry)
                page_fonts = page.get_fonts(full=True)
                if len(page_fonts) > MAX_FONT_ITEMS:
                    raise CVImportError("master CV layout font limit exceeded")
                catalog_fonts = _ordered_unique(
                    [
                        str(item[index])
                        for item in page_fonts
                        for index in (3, 4)
                        if len(item) > index and isinstance(item[index], str) and item[index]
                    ]
                )
                flags = fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES
                raw = page.get_text("dict", flags=flags)
                for raw_block in raw.get("blocks", []):
                    if raw_block.get("type") != 0:
                        continue
                    lines = raw_block.get("lines", [])
                    texts: list[str] = []
                    font_values: list[str] = []
                    size_values: list[float] = []
                    color_values: list[int] = []
                    for line in lines:
                        line_text: list[str] = []
                        for span in line.get("spans", []):
                            span_count += 1
                            if span_count > MAX_LAYOUT_SPANS:
                                raise CVImportError("master CV layout span limit exceeded")
                            line_text.append(str(span.get("text", "")))
                            font = str(span.get("font", ""))
                            if font:
                                font_values.append(font)
                            size = _finite_number(span.get("size", 0))
                            if size > 0:
                                size_values.append(round(size, 3))
                            color = span.get("color", 0)
                            if type(color) is int:
                                color_values.append(color)
                        texts.append("".join(line_text))
                    text = "\n".join(texts).strip()
                    if not text:
                        continue
                    text_chars += len(text)
                    if text_chars > MAX_LAYOUT_TEXT_CHARS:
                        raise CVImportError("master CV layout text limit exceeded")
                    bbox = tuple(_finite_number(item) for item in raw_block.get("bbox", ()))
                    if len(bbox) != 4:
                        raise CVImportError("master CV text geometry is invalid")
                    fonts = _ordered_unique(font_values) or catalog_fonts
                    sizes = _ordered_unique(size_values)
                    colors = _ordered_unique(color_values)
                    blocks.append(LayoutBlock(page_index, "text", text, bbox, fonts, sizes, colors))
                    heading = _heading(text, max(sizes, default=0.0))
                    if heading:
                        sections.append(heading)
                for image in page.get_images(full=True):
                    if image_count >= MAX_LAYOUT_IMAGES:
                        raise CVImportError("master CV layout image limit exceeded")
                    rects = page.get_image_rects(image[0])
                    for rect in rects or [fitz.Rect(0, 0, 0, 0)]:
                        bbox = tuple(float(item) for item in (rect.x0, rect.y0, rect.x1, rect.y1))
                        blocks.append(LayoutBlock(page_index, "image", "", bbox, (), (), ()))
                        image_count += 1
                if len(blocks) > MAX_LAYOUT_BLOCKS:
                    raise CVImportError("master CV layout block limit exceeded")
    except CVImportError:
        raise
    except Exception as error:
        raise CVImportError("master CV PDF layout extraction failed") from error
    return LayoutDescriptor(source_sha256, source_kind, len(pages), tuple(pages), tuple(blocks), _ordered_unique(sections), image_count)


def _sanitize_pdf_unsafe(payload: bytes) -> bytes:
    try:
        with fitz.open(stream=payload, filetype="pdf") as source:
            clean = fitz.open()
            try:
                for page_number, page in enumerate(source):
                    geometry = _geometry(page_number + 1, page.rect.width, page.rect.height)
                    target = clean.new_page(width=geometry.width_points, height=geometry.height_points)
                    target.show_pdf_page(target.rect, source, page_number, keep_proportion=False)
                clean.set_metadata({})
                result = clean.tobytes(garbage=4, clean=True, deflate=True, use_objstms=1)
            finally:
                clean.close()
        _validate_pdf_payload_unsafe(result)
        return result
    except CVImportError:
        raise
    except Exception as error:
        raise CVImportError("sanitized PDF baseline creation failed") from error


def _raster_page_png(page: Any, page_number: int) -> bytes:
    _geometry(page_number, page.rect.width, page.rect.height)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    if pixmap.width * pixmap.height > MAX_PAGE_PIXELS:
        raise CVImportError("master CV page render exceeds pixel limits")
    return pixmap.tobytes("png")


def _write_pdf_baseline_unsafe(payload: bytes, work_dir: Path) -> None:
    sanitized = _sanitize_pdf_unsafe(payload)
    atomic_write_bytes(work_dir / "baseline.pdf", sanitized)
    try:
        with fitz.open(stream=sanitized, filetype="pdf") as document:
            for page_number, page in enumerate(document, start=1):
                atomic_write_bytes(work_dir / f"page-{page_number:03d}.png", _raster_page_png(page, page_number))
    except CVImportError:
        raise
    except Exception as error:
        raise CVImportError("baseline page rendering failed") from error


def _pdf_error_category(error: Exception) -> str:
    message = str(error).casefold()
    for category in (
        "encrypted",
        "embedded content",
        "active content",
        "page count",
        "compression ratio",
        "dimensions",
        "pixel limits",
        "parser limits",
        "stream",
    ):
        if category in message:
            return category
    return "invalid"


def _pdf_worker_entry(
    input_path: str,
    result_path: str,
    mode: str,
    source_sha256: str,
    source_kind: str,
    artifact_dir: str,
    gate: Any,
) -> None:
    _apply_posix_worker_limits(MAX_WORKER_MEMORY_BYTES, MAX_RENDER_SECONDS)
    gate.wait()
    result: dict[str, Any]
    try:
        payload = _stable_single_link_bytes(Path(input_path), MAX_CV_BYTES)
        _preflight_pdf_raw(payload, check_compression=True)
        _validate_pdf_payload_unsafe(payload)
        if mode == "validate":
            result = {"ok": True}
        elif mode == "extract":
            descriptor = _pdf_layout_unsafe(payload, source_sha256, source_kind)
            _validate_descriptor_shape(descriptor)
            artifacts = path_without_link_ancestors(Path(artifact_dir))
            _write_pdf_baseline_unsafe(payload, artifacts)
            result = {"ok": True, "descriptor": _descriptor_payload(descriptor)}
        else:
            raise CVImportError("master CV PDF worker mode is invalid")
    except Exception as error:
        result = {"ok": False, "error": _pdf_error_category(error)}
    atomic_write_json(Path(result_path), result)


def _descriptor_from_json(value: object) -> LayoutDescriptor:
    if not isinstance(value, dict) or set(value) != {"source_sha256", "source_kind", "page_count", "pages", "blocks", "section_order", "image_count"}:
        raise CVImportError("PDF worker descriptor is invalid")
    pages_value, blocks_value, sections_value = value["pages"], value["blocks"], value["section_order"]
    if not isinstance(pages_value, list) or not isinstance(blocks_value, list) or not isinstance(sections_value, list):
        raise CVImportError("PDF worker descriptor is invalid")
    try:
        pages = tuple(PageGeometry(**page) for page in pages_value if isinstance(page, dict))
        blocks = tuple(
            LayoutBlock(
                page_number=block["page_number"],
                block_type=block["block_type"],
                text=block["text"],
                bbox=tuple(block["bbox"]),
                font_names=tuple(block["font_names"]),
                font_sizes=tuple(block["font_sizes"]),
                colors=tuple(block["colors"]),
                style_names=tuple(block.get("style_names", [])),
            )
            for block in blocks_value
            if isinstance(block, dict)
        )
        descriptor = LayoutDescriptor(
            source_sha256=value["source_sha256"],
            source_kind=value["source_kind"],
            page_count=value["page_count"],
            pages=pages,
            blocks=blocks,
            section_order=tuple(sections_value),
            image_count=value["image_count"],
        )
    except (KeyError, TypeError, ValueError) as error:
        raise CVImportError("PDF worker descriptor is invalid") from error
    if len(pages) != len(pages_value) or len(blocks) != len(blocks_value):
        raise CVImportError("PDF worker descriptor is invalid")
    _validate_descriptor_shape(descriptor)
    return descriptor


def _run_pdf_worker(
    payload: bytes,
    *,
    mode: str,
    source_sha256: str | None = None,
    source_kind: str | None = None,
    output_dir: Path | None = None,
) -> LayoutDescriptor | None:
    if mode not in {"validate", "extract"}:
        raise CVImportError("PDF worker mode is invalid")
    parent = _safe_work_dir(output_dir) if output_dir is not None else None
    with tempfile.TemporaryDirectory(prefix=".cv-pdf-worker-", dir=parent) as temporary_name:
        temporary = path_without_link_ancestors(Path(temporary_name))
        input_path = temporary / "input.pdf"
        result_path = temporary / "result.json"
        artifacts = temporary / "artifacts"
        artifacts.mkdir()
        atomic_write_bytes(input_path, payload)
        succeeded = _run_isolated_worker(
            _pdf_worker_entry,
            (
                str(input_path),
                str(result_path),
                mode,
                source_sha256 or "0" * 64,
                source_kind or "pdf",
                str(artifacts),
            ),
            timeout_seconds=MAX_RENDER_SECONDS,
            memory_bytes=MAX_WORKER_MEMORY_BYTES,
        )
        if not succeeded:
            raise CVImportError("master CV PDF isolated worker failed")
        try:
            result_bytes = _stable_single_link_bytes(result_path, MAX_WORKER_IPC_BYTES)
            result = json.loads(result_bytes.decode("utf-8"))
        except (PolicyError, UnicodeDecodeError, json.JSONDecodeError, OSError, ValueError) as error:
            raise CVImportError("master CV PDF isolated worker returned invalid output") from error
        if not isinstance(result, dict) or type(result.get("ok")) is not bool:
            raise CVImportError("master CV PDF isolated worker returned invalid output")
        if not result["ok"]:
            category = result.get("error")
            if not isinstance(category, str) or category not in {"encrypted", "embedded content", "active content", "page count", "compression ratio", "dimensions", "pixel limits", "parser limits", "stream", "invalid"}:
                category = "invalid"
            raise CVImportError(f"master CV PDF {category}")
        if mode == "validate":
            if set(result) != {"ok"}:
                raise CVImportError("master CV PDF isolated worker returned invalid output")
            return None
        if set(result) != {"ok", "descriptor"} or parent is None:
            raise CVImportError("master CV PDF isolated worker returned invalid output")
        descriptor = _descriptor_from_json(result["descriptor"])
        for page_number in range(1, descriptor.page_count + 1):
            source_image = artifacts / f"page-{page_number:03d}.png"
            try:
                image_payload = _stable_single_link_bytes(source_image, MAX_WORKER_ARTIFACT_BYTES)
            except PolicyError as error:
                raise CVImportError("master CV PDF isolated worker artifact is invalid") from error
            if not image_payload.startswith(b"\x89PNG\r\n\x1a\n"):
                raise CVImportError("master CV PDF isolated worker artifact is invalid")
            atomic_write_bytes(parent / source_image.name, image_payload)
        try:
            baseline = _stable_single_link_bytes(artifacts / "baseline.pdf", MAX_CV_BYTES)
        except PolicyError as error:
            raise CVImportError("master CV PDF isolated worker artifact is invalid") from error
        atomic_write_bytes(parent / "baseline.pdf", baseline)
        return descriptor


def _validate_pdf_payload(payload: bytes) -> None:
    _preflight_pdf_raw(payload, check_compression=False)
    _run_pdf_worker(payload, mode="validate")


def _extract_pdf_isolated(payload: bytes, source_sha256: str, source_kind: str, output_dir: Path) -> LayoutDescriptor:
    _preflight_pdf_raw(payload, check_compression=False)
    descriptor = _run_pdf_worker(
        payload,
        mode="extract",
        source_sha256=source_sha256,
        source_kind=source_kind,
        output_dir=output_dir,
    )
    if descriptor is None:
        raise CVImportError("master CV PDF isolated worker returned invalid output")
    return descriptor


def _find_libreoffice() -> Path | None:
    for name in ("soffice.exe", "soffice"):
        found = shutil.which(name)
        if found:
            try:
                return path_without_link_ancestors(Path(found))
            except (OSError, ValueError):
                continue
    candidates: list[Path] = []
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
        root = os.environ.get(variable)
        if root:
            candidates.append(Path(root) / "LibreOffice" / "program" / "soffice.exe")
    for candidate in candidates:
        try:
            checked = path_without_link_ancestors(candidate)
            if checked.is_file() and not is_symlink_or_reparse(checked):
                return checked
        except (OSError, ValueError):
            continue
    return None


def _renderer_environment(executable: Path, isolation: Path) -> dict[str, str]:
    environment = {
        "PATH": str(executable.parent),
        "TEMP": str(isolation),
        "TMP": str(isolation),
        "HOME": str(isolation),
        "USERPROFILE": str(isolation),
        "APPDATA": str(isolation),
        "LOCALAPPDATA": str(isolation),
        "SAL_DISABLE_OPENCL": "1",
        "SAL_DISABLEGL": "1",
    }
    system_root = os.environ.get("SystemRoot")
    if system_root:
        environment["SystemRoot"] = system_root
    return environment


def _libreoffice_worker(
    executable_path: str,
    docx_path: str,
    output_path: str,
    isolation_path: str,
    gate: Any,
) -> None:
    _apply_posix_worker_limits(MAX_WORKER_MEMORY_BYTES, MAX_RENDER_SECONDS)
    gate.wait()
    executable = path_without_link_ancestors(Path(executable_path))
    source = path_without_link_ancestors(Path(docx_path))
    output = path_without_link_ancestors(Path(output_path))
    isolation = path_without_link_ancestors(Path(isolation_path))
    profile = isolation / "profile"
    profile.mkdir()
    command = [
        str(executable),
        "--headless",
        "--invisible",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        "--norestore",
        "--nofirststartwizard",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(isolation),
        str(source),
    ]
    try:
        result = subprocess.run(
            command,
            cwd=isolation,
            env=_renderer_environment(executable, isolation),
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=MAX_RENDER_SECONDS,
            check=False,
            shell=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.SubprocessError):
        raise SystemExit(70)
    generated = isolation / f"{source.stem}.pdf"
    if result.returncode != 0 or not generated.is_file() or generated != output:
        raise SystemExit(70)


def _render_docx_with_libreoffice(executable: Path, docx_path: Path, work_dir: Path) -> bytes | None:
    invocation = work_dir / f"libreoffice-{uuid.uuid4().hex}"
    try:
        invocation = confined_mkdir(work_dir, invocation)
        private_source = invocation / "input.docx"
        output = invocation / "input.pdf"
        source_payload = _stable_single_link_bytes(docx_path, MAX_CV_BYTES)
        atomic_write_bytes(private_source, source_payload)
        succeeded = _run_isolated_worker(
            _libreoffice_worker,
            (str(executable), str(private_source), str(output), str(invocation)),
            timeout_seconds=MAX_RENDER_SECONDS,
            memory_bytes=MAX_WORKER_MEMORY_BYTES,
        )
        if not succeeded or not output.is_file():
            return None
        return _stable_single_link_bytes(output, MAX_CV_BYTES)
    except (ConfinementError, OSError, PolicyError, ValueError):
        return None
    finally:
        try:
            checked = confined_path(work_dir, invocation)
            if checked.exists() and checked.is_dir() and not is_symlink_or_reparse(checked):
                shutil.rmtree(checked)
        except (ConfinementError, OSError, ValueError):
            pass


def _apply_posix_worker_limits(memory_bytes: int, timeout_seconds: int) -> None:
    if os.name == "nt":
        return
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (memory_bytes, memory_bytes))
        resource.setrlimit(resource.RLIMIT_CPU, (timeout_seconds, timeout_seconds + 1))
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    except (ImportError, OSError, ValueError):
        raise SystemExit(70)


def _windows_job_for_process(process_id: int, memory_bytes: int) -> int | None:
    if os.name != "nt":
        return None
    import ctypes
    from ctypes import wintypes

    class IO_COUNTERS(ctypes.Structure):
        _fields_ = [
            ("ReadOperationCount", ctypes.c_ulonglong),
            ("WriteOperationCount", ctypes.c_ulonglong),
            ("OtherOperationCount", ctypes.c_ulonglong),
            ("ReadTransferCount", ctypes.c_ulonglong),
            ("WriteTransferCount", ctypes.c_ulonglong),
            ("OtherTransferCount", ctypes.c_ulonglong),
        ]

    class BASIC_LIMITS(ctypes.Structure):
        _fields_ = [
            ("PerProcessUserTimeLimit", ctypes.c_longlong),
            ("PerJobUserTimeLimit", ctypes.c_longlong),
            ("LimitFlags", wintypes.DWORD),
            ("MinimumWorkingSetSize", ctypes.c_size_t),
            ("MaximumWorkingSetSize", ctypes.c_size_t),
            ("ActiveProcessLimit", wintypes.DWORD),
            ("Affinity", ctypes.c_size_t),
            ("PriorityClass", wintypes.DWORD),
            ("SchedulingClass", wintypes.DWORD),
        ]

    class EXTENDED_LIMITS(ctypes.Structure):
        _fields_ = [
            ("BasicLimitInformation", BASIC_LIMITS),
            ("IoInfo", IO_COUNTERS),
            ("ProcessMemoryLimit", ctypes.c_size_t),
            ("JobMemoryLimit", ctypes.c_size_t),
            ("PeakProcessMemoryUsed", ctypes.c_size_t),
            ("PeakJobMemoryUsed", ctypes.c_size_t),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateJobObjectW.restype = wintypes.HANDLE
    kernel32.OpenProcess.restype = wintypes.HANDLE
    job = kernel32.CreateJobObjectW(None, None)
    if not job:
        return None
    information = EXTENDED_LIMITS()
    information.BasicLimitInformation.LimitFlags = 0x2000 | 0x0200
    information.JobMemoryLimit = memory_bytes
    if not kernel32.SetInformationJobObject(job, 9, ctypes.byref(information), ctypes.sizeof(information)):
        kernel32.CloseHandle(job)
        return None
    process = kernel32.OpenProcess(0x0001 | 0x0100 | 0x0400, False, process_id)
    if not process:
        kernel32.CloseHandle(job)
        return None
    try:
        if not kernel32.AssignProcessToJobObject(job, process):
            kernel32.CloseHandle(job)
            return None
    finally:
        kernel32.CloseHandle(process)
    return int(job)


def _terminate_worker_tree(process: Any, job: int | None) -> None:
    if os.name == "nt" and job:
        try:
            import ctypes

            ctypes.WinDLL("kernel32", use_last_error=True).TerminateJobObject(job, 70)
        except BaseException:
            pass
    try:
        alive = process.is_alive()
    except BaseException:
        alive = True
    if alive:
        try:
            process.terminate()
        except BaseException:
            pass
    try:
        process.join(timeout=5)
    except BaseException:
        pass
    try:
        alive = process.is_alive()
    except BaseException:
        alive = True
    if alive and hasattr(process, "kill"):
        try:
            process.kill()
        except BaseException:
            pass
        try:
            process.join(timeout=5)
        except BaseException:
            pass


def _close_worker_job(job: int | None) -> None:
    if os.name == "nt" and job:
        import ctypes

        ctypes.WinDLL("kernel32", use_last_error=True).CloseHandle(job)


def _isolated_worker_entry(target: Any, arguments: tuple[Any, ...], gate: Any) -> None:
    stdout_descriptor: int | None = None
    stderr_descriptor: int | None = None
    try:
        try:
            stdout_descriptor = os.open(os.devnull, os.O_WRONLY)
            stderr_descriptor = os.open(os.devnull, os.O_WRONLY)
            os.dup2(stdout_descriptor, 1)
            os.dup2(stderr_descriptor, 2)
        finally:
            for descriptor in (stdout_descriptor, stderr_descriptor):
                if descriptor is not None and descriptor not in {1, 2}:
                    try:
                        os.close(descriptor)
                    except OSError:
                        pass
        stdout_sink = open(os.devnull, "w", encoding="utf-8")
        stderr_sink = open(os.devnull, "w", encoding="utf-8")
        sys.stdout = stdout_sink
        sys.stderr = stderr_sink
        try:
            target(*arguments, gate)
        except BaseException:
            raise SystemExit(70) from None
    except BaseException:
        raise SystemExit(70) from None


def _run_isolated_worker(target: Any, arguments: tuple[Any, ...], *, timeout_seconds: int, memory_bytes: int) -> bool:
    try:
        context = multiprocessing.get_context("spawn")
        gate = context.Event()
        process = context.Process(target=_isolated_worker_entry, args=(target, arguments, gate), daemon=False)
    except BaseException:
        return False
    job: int | None = None
    start_attempted = False
    succeeded = False
    try:
        start_attempted = True
        process.start()
        if os.name == "nt":
            process_id = process.pid
            if type(process_id) is not int or process_id <= 0:
                raise RuntimeError("isolated worker process ID is unavailable")
            job = _windows_job_for_process(process_id, memory_bytes)
            if job is None:
                raise RuntimeError("isolated worker Job setup failed")
        gate.set()
        process.join(timeout_seconds)
        if process.is_alive():
            raise TimeoutError("isolated worker timed out")
        succeeded = process.exitcode == 0
    except BaseException:
        succeeded = False
    finally:
        if start_attempted and not succeeded:
            try:
                _terminate_worker_tree(process, job)
            except BaseException:
                pass
        try:
            _close_worker_job(job)
        except BaseException:
            succeeded = False
    return succeeded


def _render_docx_with_word(_docx_path: Path, _output_path: Path, _work_dir: Path) -> None:
    # Word COM has no reliable non-admin OS network-denial boundary, so this route is disabled.
    return None


def _render_docx_to_pdf(payload: bytes, work_dir: Path) -> bytes:
    token = uuid.uuid4().hex
    source = work_dir / f"source-{token}.docx"
    word_output = work_dir / f"word-{token}.pdf"
    atomic_write_bytes(source, payload)
    try:
        executable = _find_libreoffice()
        converted = _render_docx_with_libreoffice(executable, source, work_dir) if executable is not None else None
        if converted is None:
            converted = _render_docx_with_word(source, word_output, work_dir)
        if converted is None:
            raise CVImportError("DOCX visual import requires an available isolated LibreOffice renderer; export the master CV as PDF and retry")
        _validate_pdf_payload(converted)
        return converted
    finally:
        for path in (source, work_dir / f"{source.stem}.pdf", word_output):
            try:
                if path.exists() and not is_symlink_or_reparse(path):
                    path.unlink()
            except OSError:
                pass


def _run_properties(run: Any, paragraph: Any) -> tuple[tuple[str, ...], tuple[float, ...], tuple[int, ...]]:
    names: list[str] = []
    sizes: list[float] = []
    colors: list[int] = []
    for font in (getattr(run, "font", None), getattr(getattr(paragraph, "style", None), "font", None)):
        if font is None:
            continue
        if font.name:
            names.append(str(font.name))
        if font.size:
            sizes.append(round(float(font.size.pt), 3))
        rgb = getattr(getattr(font, "color", None), "rgb", None)
        if rgb is not None:
            colors.append(int(str(rgb), 16))
    return _ordered_unique(names), _ordered_unique(sizes), _ordered_unique(colors)


def _docx_semantics(payload: bytes, visual: LayoutDescriptor) -> LayoutDescriptor:
    try:
        document = Document(io.BytesIO(payload))
    except Exception as error:
        raise CVImportError("master CV DOCX semantic extraction failed") from error
    blocks = list(visual.blocks)
    sections = list(visual.section_order)
    first_geometry = visual.pages[0]
    page_number = 1

    def append_paragraph(paragraph: Any, block_type: str) -> None:
        nonlocal blocks
        text = paragraph.text.strip()
        if not text:
            return
        names: list[str] = []
        sizes: list[float] = []
        colors: list[int] = []
        for run in paragraph.runs:
            run_names, run_sizes, run_colors = _run_properties(run, paragraph)
            names.extend(run_names)
            sizes.extend(run_sizes)
            colors.extend(run_colors)
        style = str(paragraph.style.name) if paragraph.style is not None and paragraph.style.name else ""
        blocks.append(
            LayoutBlock(
                page_number,
                block_type,
                text,
                (0.0, 0.0, 0.0, 0.0),
                _ordered_unique(names),
                _ordered_unique(sizes),
                _ordered_unique(colors),
                (style,) if style else (),
            )
        )
        heading = _heading(text, max(sizes, default=0.0))
        if "heading" in style.casefold() or heading:
            sections.append(" ".join(text.split()))

    for paragraph in document.paragraphs:
        append_paragraph(paragraph, "paragraph")
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
            if text:
                blocks.append(LayoutBlock(page_number, "table", text, (0.0, 0.0, 0.0, 0.0), (), (), (), (str(table.style.name),) if table.style is not None and table.style.name else ()))
    pages = list(visual.pages)
    for index, section in enumerate(document.sections):
        margins = (
            float(section.top_margin.pt) if section.top_margin is not None else 0.0,
            float(section.right_margin.pt) if section.right_margin is not None else 0.0,
            float(section.bottom_margin.pt) if section.bottom_margin is not None else 0.0,
            float(section.left_margin.pt) if section.left_margin is not None else 0.0,
        )
        if index < len(pages):
            pages[index] = replace(
                pages[index],
                margin_top_points=margins[0],
                margin_right_points=margins[1],
                margin_bottom_points=margins[2],
                margin_left_points=margins[3],
            )
        elif index == 0:
            pages[0] = replace(first_geometry, margin_top_points=margins[0], margin_right_points=margins[1], margin_bottom_points=margins[2], margin_left_points=margins[3])
        blocks.append(LayoutBlock(min(index + 1, visual.page_count), "section", f"section {index + 1}", (0.0, 0.0, 0.0, 0.0), (), (), (), ()))
        for paragraph in section.header.paragraphs:
            append_paragraph(paragraph, "header")
        for paragraph in section.footer.paragraphs:
            append_paragraph(paragraph, "footer")
    semantic_images = len(document.inline_shapes)
    image_count = visual.image_count
    while image_count < semantic_images:
        blocks.append(LayoutBlock(1, "image", "", (0.0, 0.0, 0.0, 0.0), (), (), (), ()))
        image_count += 1
    if len(blocks) > MAX_LAYOUT_BLOCKS or sum(len(block.text) for block in blocks) > MAX_LAYOUT_TEXT_CHARS:
        raise CVImportError("master CV DOCX layout limits exceeded")
    return LayoutDescriptor(visual.source_sha256, "docx", visual.page_count, tuple(pages), tuple(blocks), _ordered_unique(sections), image_count)


def _extract_layout_to_work_dir(source: SourceCV, work_dir: Path) -> LayoutDescriptor:
    payload = _current_payload(source)
    destination = _safe_work_dir(work_dir)
    if source.kind == "pdf":
        descriptor = _extract_pdf_isolated(payload, source.sha256, "pdf", destination)
    else:
        visual_payload = _render_docx_to_pdf(payload, destination)
        _validate_pdf_payload(visual_payload)
        visual = _extract_pdf_isolated(visual_payload, source.sha256, "docx", destination)
        descriptor = _docx_semantics(payload, visual)
    _validate_descriptor_shape(descriptor)
    return descriptor


def extract_layout(source: SourceCV, project_root: Path) -> LayoutDescriptor:
    try:
        root = confined_path(project_root, project_root)
        with exclusive_project_lock(root, name=".cv-import.lock"):
            raw_path, registered = _registered_source(root)
            if _source_identity(source) != _source_identity(registered):
                raise CVImportError("source is not the registered master CV")
            destination = confined_mkdir(root, root / "private" / "cv-creator" / "baseline")
            descriptor = _extract_layout_to_work_dir(registered, destination)
            after_path, after = _registered_source(root)
            if after_path != raw_path or _source_identity(after) != _source_identity(registered):
                raise CVImportError("registered master CV changed during extraction")
            token = uuid.uuid4().hex
            descriptor = replace(descriptor, _attestation=token)
            attestation = _ImportAttestation(
                project_root=root,
                source_path=raw_path,
                source_identity=_source_identity(after),
                descriptor_digest=_descriptor_digest(descriptor),
            )
            with _ATTESTATION_LOCK:
                for stale_token, stale in tuple(_ATTESTATIONS.items()):
                    if stale.project_root == root:
                        del _ATTESTATIONS[stale_token]
                _ATTESTATIONS[token] = attestation
            return descriptor
    except CVImportError:
        raise
    except (ConfinementError, OSError, ValueError) as error:
        raise CVImportError("registered master CV baseline is unavailable") from error


def _validate_descriptor_shape(descriptor: LayoutDescriptor) -> None:
    if type(descriptor) is not LayoutDescriptor:
        raise CVImportError("layout descriptor type is invalid")
    if SHA256.fullmatch(descriptor.source_sha256) is None or descriptor.source_kind not in {"pdf", "docx"}:
        raise CVImportError("layout descriptor source is invalid")
    if type(descriptor.page_count) is not int or not 1 <= descriptor.page_count <= MAX_CV_PAGES or len(descriptor.pages) != descriptor.page_count:
        raise CVImportError("layout descriptor page total is invalid")
    page_map: dict[int, PageGeometry] = {}
    for expected, page in enumerate(descriptor.pages, start=1):
        if type(page) is not PageGeometry or type(page.page_number) is not int or page.page_number != expected:
            raise CVImportError("layout descriptor page numbering is invalid")
        geometry = _geometry(page.page_number, page.width_points, page.height_points)
        for margin in (page.margin_top_points, page.margin_right_points, page.margin_bottom_points, page.margin_left_points):
            if margin is not None and not 0 <= _finite_number(margin) <= max(geometry.width_points, geometry.height_points):
                raise CVImportError("layout descriptor margins are invalid")
        horizontal = (page.margin_left_points or 0.0) + (page.margin_right_points or 0.0)
        vertical = (page.margin_top_points or 0.0) + (page.margin_bottom_points or 0.0)
        if horizontal > geometry.width_points or vertical > geometry.height_points:
            raise CVImportError("layout descriptor margins are invalid")
        page_map[page.page_number] = geometry
    if type(descriptor.blocks) is not tuple or len(descriptor.blocks) > MAX_LAYOUT_BLOCKS:
        raise CVImportError("layout descriptor block total is invalid")
    image_count = text_total = 0
    for block in descriptor.blocks:
        if type(block) is not LayoutBlock or type(block.page_number) is not int or block.page_number not in page_map or block.block_type not in ALLOWED_BLOCK_TYPES:
            raise CVImportError("layout descriptor block is invalid")
        if not isinstance(block.text, str) or len(block.text) > MAX_LAYOUT_TEXT_CHARS:
            raise CVImportError("layout descriptor text is invalid")
        text_total += len(block.text)
        if text_total > MAX_LAYOUT_TEXT_CHARS:
            raise CVImportError("layout descriptor text total is invalid")
        if type(block.bbox) is not tuple or len(block.bbox) != 4:
            raise CVImportError("layout descriptor block geometry is invalid")
        x0, y0, x1, y1 = (_finite_number(value) for value in block.bbox)
        geometry = page_map[block.page_number]
        if not (0 <= x0 <= x1 <= geometry.width_points + 1 and 0 <= y0 <= y1 <= geometry.height_points + 1):
            raise CVImportError("layout descriptor block geometry is invalid")
        if (
            type(block.font_names) is not tuple
            or type(block.font_sizes) is not tuple
            or type(block.colors) is not tuple
            or type(block.style_names) is not tuple
            or len(block.font_names) > MAX_FONT_ITEMS
            or len(block.font_sizes) > MAX_FONT_ITEMS
            or len(block.colors) > MAX_FONT_ITEMS
            or len(block.style_names) > MAX_STYLE_ITEMS
            or any(not isinstance(name, str) or not name or len(name) > 256 for name in block.font_names + block.style_names)
            or any(not 0 < _finite_number(size) <= 1_000 for size in block.font_sizes)
            or any(type(color) is not int or not 0 <= color <= 0xFFFFFF for color in block.colors)
        ):
            raise CVImportError("layout descriptor font metadata is invalid")
        if block.block_type == "image":
            image_count += 1
    if type(descriptor.image_count) is not int or descriptor.image_count != image_count or image_count > MAX_LAYOUT_IMAGES:
        raise CVImportError("layout descriptor image total is invalid")
    if (
        type(descriptor.section_order) is not tuple
        or len(descriptor.section_order) > MAX_SECTION_ITEMS
        or any(not isinstance(item, str) or not item.strip() or len(item) > 256 for item in descriptor.section_order)
    ):
        raise CVImportError("layout descriptor section order is invalid")


def _registered_source(project_root: Path) -> tuple[str, SourceCV]:
    try:
        root = confined_path(project_root, project_root)
        runtime = load_runtime(root)
        raw_path = runtime["master_cv_path"]
        digest = runtime["master_cv_sha256"]
        if raw_path not in {"private/master-cv.pdf", "private/master-cv.docx"} or SHA256.fullmatch(digest) is None:
            raise CVImportError("registered master CV is unavailable")
        relative = PurePosixPath(raw_path)
        source = validate_source_cv(confined_path(root, root.joinpath(*relative.parts)))
        if source.sha256 != digest:
            raise CVImportError("registered master CV hash has changed")
        return raw_path, source
    except CVImportError:
        raise
    except (KeyError, OnboardingError, OSError, ValueError) as error:
        raise CVImportError("registered master CV is unavailable") from error


def write_layout_descriptor(project_root: Path, descriptor: LayoutDescriptor) -> Path:
    _validate_descriptor_shape(descriptor)
    try:
        root = confined_path(project_root, project_root)
        token = descriptor._attestation
        if not isinstance(token, str) or not token:
            raise CVImportError("layout descriptor has no current import attestation")
        with exclusive_project_lock(root, name=".cv-import.lock"):
            with _ATTESTATION_LOCK:
                attestation = _ATTESTATIONS.get(token)
                if (
                    attestation is None
                    or attestation.project_root != root
                    or attestation.descriptor_digest != _descriptor_digest(descriptor)
                ):
                    raise CVImportError("layout descriptor does not match current import attestation")
                del _ATTESTATIONS[token]
            raw_path, source = _registered_source(root)
            if (
                raw_path != attestation.source_path
                or _source_identity(source) != attestation.source_identity
                or descriptor.source_sha256 != source.sha256
                or descriptor.source_kind != source.kind
            ):
                raise CVImportError("layout descriptor does not match registered master CV")
            creator = confined_mkdir(root, root / "private" / "cv-creator")
            destination = confined_path(root, creator / "layout.json")
            prior: bytes | None = None
            if destination.exists() or is_symlink_or_reparse(destination):
                metadata = destination.lstat()
                if is_symlink_or_reparse(destination) or not stat.S_ISREG(metadata.st_mode) or metadata.st_nlink != 1:
                    raise CVImportError("layout destination is unsafe")
                try:
                    prior = _stable_single_link_bytes(destination, MAX_WORKER_IPC_BYTES)
                except PolicyError as error:
                    raise CVImportError("layout destination is unsafe") from error
            payload = _descriptor_payload(descriptor)
            payload["source_path"] = raw_path
            published = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
            atomic_write_json(destination, payload, project_root=root)
            try:
                after_path, after = _registered_source(root)
                if after_path != raw_path or _source_identity(after) != attestation.source_identity:
                    raise CVImportError("registered master CV changed during layout commit")
                committed = _stable_single_link_bytes(destination, MAX_WORKER_IPC_BYTES)
                if committed != published:
                    raise CVImportError("layout destination changed during layout commit")
            except Exception as error:
                try:
                    current = _stable_single_link_bytes(destination, MAX_WORKER_IPC_BYTES)
                    if current == published:
                        if prior is None:
                            destination.unlink()
                        else:
                            atomic_write_bytes(destination, prior, project_root=root)
                except (OSError, PolicyError, ValueError):
                    pass
                raise CVImportError("registered master CV changed during layout commit") from error
            return destination
    except CVImportError:
        raise
    except (ConfinementError, OSError, ValueError) as error:
        raise CVImportError("layout destination is unsafe") from error


def main() -> int:
    parser = argparse.ArgumentParser(description="Import registered private master CV layout.")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    arguments = parser.parse_args()
    try:
        _, source = _registered_source(arguments.root)
        root = confined_path(arguments.root, arguments.root)
        descriptor = extract_layout(source, root)
        write_layout_descriptor(root, descriptor)
        output = {"ok": True, "kind": source.kind, "page_count": descriptor.page_count}
        code = 0
    except Exception:
        output = {"ok": False, "error": "CV import failed"}
        code = 2
    print(json.dumps(output, ensure_ascii=False))
    return code


if __name__ == "__main__":
    raise SystemExit(main())
