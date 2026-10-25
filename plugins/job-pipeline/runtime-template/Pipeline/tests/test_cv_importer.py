from __future__ import annotations

import hashlib
import json
import os
import struct
import subprocess
import sys
import types
import zipfile
import zlib
from contextlib import contextmanager
from dataclasses import FrozenInstanceError, replace
from pathlib import Path
from xml.etree import ElementTree

import fitz
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from Pipeline.cv_importer import (
    CVImportError,
    MAX_CV_BYTES,
    MAX_CV_PAGES,
    _extract_layout_to_work_dir,
    extract_layout,
    validate_source_cv,
    write_layout_descriptor,
)
from Pipeline.initialize import initialize_project
from Pipeline.onboarding import register_master_cv


PNG_1X1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c49444154789c63606060000000040001f61738550000000049454e44ae426082"
)


def leaking_worker_target(secret: str, gate) -> None:
    gate.wait()
    print(secret, flush=True)
    sys.stderr.write(secret + "\n")
    raise RuntimeError(secret)


def make_pdf(
    path: Path,
    *,
    pages: int = 1,
    link: bool = False,
    metadata: bool = False,
) -> bytes:
    document = fitz.open()
    for number in range(pages):
        page = document.new_page(width=595, height=842)
        page.insert_text((48, 60), "SYNTHETIC CV", fontsize=18, fontname="helv")
        page.insert_text((48, 100), "EXPERIENCE", fontsize=11, color=(0.1, 0.2, 0.3))
        page.insert_text((48, 122), f"Verified synthetic statement {number + 1}", fontsize=9)
        page.insert_image(fitz.Rect(500, 40, 520, 60), stream=PNG_1X1)
        if link:
            page.insert_link(
                {
                    "kind": fitz.LINK_URI,
                    "from": fitz.Rect(48, 130, 180, 145),
                    "uri": "https://example.test/profile",
                }
            )
    if metadata:
        document.set_metadata({"author": "SYNTHETIC PRIVATE AUTHOR", "title": "Synthetic"})
    document.save(path, deflate=True, use_objstms=1)
    document.close()
    return path.read_bytes()


def make_docx(path: Path) -> bytes:
    document = Document()
    style = document.styles.add_style("Synthetic Heading", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(15)
    heading = document.add_paragraph(style=style)
    run = heading.add_run("SYNTHETIC PROFILE")
    run.bold = True
    run.font.color.rgb = RGBColor(12, 34, 56)
    document.add_paragraph("EXPERIENCE")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Synthetic role"
    table.cell(0, 1).text = "2030"
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.8)
    section.header.paragraphs[0].text = "Synthetic header"
    section.footer.paragraphs[0].text = "Synthetic footer"
    image = path.with_suffix(".png")
    image.write_bytes(PNG_1X1)
    document.add_picture(str(image), width=Inches(0.1))
    document.save(path)
    image.unlink()
    return path.read_bytes()


def rewrite_member(path: Path, name: str, payload: bytes) -> None:
    temporary = path.with_suffix(path.suffix + ".rewrite")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, payload if item.filename == name else source.read(item))
    temporary.replace(path)


def rename_member(path: Path, old_name: str, new_name: str) -> None:
    temporary = path.with_suffix(path.suffix + ".rename")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(new_name if item.filename == old_name else item, source.read(item))
    temporary.replace(path)


def add_external_relationship(path: Path, *, relation_type: str, target: str) -> None:
    name = "word/_rels/document.xml.rels"
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read(name))
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    ElementTree.SubElement(
        root,
        f"{{{namespace}}}Relationship",
        {
            "Id": "rIdSyntheticExternal",
            "Type": relation_type,
            "Target": target,
            "TargetMode": "External",
        },
    )
    rewrite_member(path, name, ElementTree.tostring(root, encoding="utf-8", xml_declaration=True))


def add_internal_relationship(path: Path, *, relation_type: str, target: str) -> None:
    name = "word/_rels/document.xml.rels"
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read(name))
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    ElementTree.SubElement(
        root,
        f"{{{namespace}}}Relationship",
        {"Id": "rIdSyntheticInternal", "Type": relation_type, "Target": target},
    )
    rewrite_member(path, name, ElementTree.tostring(root, encoding="utf-8", xml_declaration=True))


def png_with_dimensions(width: int, height: int) -> bytes:
    signature = b"\x89PNG\r\n\x1a\n"
    data = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    chunk = b"IHDR" + data
    return signature + struct.pack(">I", len(data)) + chunk + struct.pack(">I", zlib.crc32(chunk) & 0xFFFFFFFF) + b"\x00\x00\x00\x00IEND\xaeB`\x82"


def fake_docx_renderer(_payload: bytes, work_dir: Path) -> bytes:
    output = work_dir / "converted.pdf"
    return make_pdf(output)


def test_pdf_import_extracts_geometry_text_fonts_colors_images_and_order(tmp_path: Path):
    path = tmp_path / "master.pdf"
    make_pdf(path)
    source = validate_source_cv(path)
    descriptor = _extract_layout_to_work_dir(source, tmp_path / "work")
    assert source.kind == "pdf"
    assert descriptor.page_count == 1
    assert descriptor.pages[0].width_points == pytest.approx(595, abs=1)
    assert any("SYNTHETIC CV" in block.text for block in descriptor.blocks)
    assert any(block.font_names and block.font_sizes and block.colors for block in descriptor.blocks if block.block_type == "text")
    assert descriptor.image_count == 1
    assert descriptor.section_order == ("SYNTHETIC CV", "EXPERIENCE")
    assert descriptor.source_sha256 == source.sha256


def test_source_identity_is_frozen_and_matches_exact_file_bytes(tmp_path: Path):
    path = tmp_path / "master.pdf"
    payload = make_pdf(path)
    source = validate_source_cv(path)
    assert source.sha256 == hashlib.sha256(payload).hexdigest()
    assert source.size_bytes == len(payload)
    with pytest.raises(FrozenInstanceError):
        source.sha256 = "0" * 64  # type: ignore[misc]


@pytest.mark.parametrize("suffix,payload", [(".txt", b"plain"), (".pdf", b"PK\x03\x04not-pdf"), (".docx", b"%PDF-1.7")])
def test_import_rejects_wrong_extension_or_magic(tmp_path: Path, suffix: str, payload: bytes):
    path = tmp_path / f"master{suffix}"
    path.write_bytes(payload)
    with pytest.raises(CVImportError):
        validate_source_cv(path)


def test_import_rejects_empty_oversize_directory_and_hardlink(tmp_path: Path):
    empty = tmp_path / "empty.pdf"
    empty.write_bytes(b"")
    with pytest.raises(CVImportError, match="size"):
        validate_source_cv(empty)
    large = tmp_path / "large.pdf"
    large.write_bytes(b"0" * (MAX_CV_BYTES + 1))
    with pytest.raises(CVImportError, match="size"):
        validate_source_cv(large)
    directory = tmp_path / "directory.pdf"
    directory.mkdir()
    with pytest.raises(CVImportError, match="regular"):
        validate_source_cv(directory)
    original = tmp_path / "original.pdf"
    make_pdf(original)
    linked = tmp_path / "linked.pdf"
    os.link(original, linked)
    with pytest.raises(CVImportError, match="hardlink"):
        validate_source_cv(linked)


def test_import_rejects_symlink_ancestor(tmp_path: Path):
    outside = tmp_path / "outside"
    outside.mkdir()
    make_pdf(outside / "master.pdf")
    linked = tmp_path / "linked"
    try:
        linked.symlink_to(outside, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"symlink unavailable: {error}")
    with pytest.raises(CVImportError, match="symlink or reparse"):
        validate_source_cv(linked / "master.pdf")


def test_extract_rejects_post_validation_mutation(tmp_path: Path):
    path = tmp_path / "master.pdf"
    make_pdf(path)
    source = validate_source_cv(path)
    make_pdf(path.with_name("replacement.pdf"), pages=2)
    path.write_bytes(path.with_name("replacement.pdf").read_bytes())
    with pytest.raises(CVImportError, match="changed after validation"):
        _extract_layout_to_work_dir(source, tmp_path / "work")


def test_import_rejects_encrypted_active_embedded_and_too_many_page_pdfs(tmp_path: Path):
    encrypted = tmp_path / "encrypted.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "Synthetic")
    document.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    document.close()
    with pytest.raises(CVImportError, match="encrypted"):
        validate_source_cv(encrypted)

    active = tmp_path / "active.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Synthetic")
    document.xref_set_key(page.xref, "AA", "<</O<</S/JavaScript/JS(app.alert('x'))>>>>")
    document.save(active, deflate=True, garbage=4, use_objstms=1)
    document.close()
    with pytest.raises(CVImportError, match="active content"):
        validate_source_cv(active)

    embedded = tmp_path / "embedded.pdf"
    document = fitz.open()
    document.new_page()
    document.embfile_add("payload.bin", b"synthetic")
    document.save(embedded)
    document.close()
    with pytest.raises(CVImportError, match="embedded"):
        validate_source_cv(embedded)

    many = tmp_path / "many.pdf"
    make_pdf(many, pages=MAX_CV_PAGES + 1)
    with pytest.raises(CVImportError, match="page count"):
        validate_source_cv(many)


def test_import_rejects_javascript_uri_action(tmp_path: Path):
    path = tmp_path / "javascript-uri.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Synthetic")
    page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(72, 80, 180, 100), "uri": "javascript:alert(1)"})
    document.save(path, deflate=True, use_objstms=1)
    document.close()
    with pytest.raises(CVImportError, match="active content"):
        validate_source_cv(path)


def test_pdf_validation_delegates_parse_to_isolated_worker(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    path = tmp_path / "master.pdf"
    make_pdf(path)
    calls: list[str] = []

    def fake_worker(payload, *, mode, source_sha256=None, source_kind=None, output_dir=None):
        assert payload.startswith(b"%PDF-")
        calls.append(mode)
        return None

    monkeypatch.setattr(cv_importer, "_run_pdf_worker", fake_worker, raising=False)
    assert validate_source_cv(path).kind == "pdf"
    assert calls == ["validate"]


def test_compressed_stream_bomb_rejected_before_pymupdf_parse(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    compressed = zlib.compress(b"A" * 2_000_000, level=9)
    payload = (
        b"%PDF-1.7\n1 0 obj\n<</Length "
        + str(len(compressed)).encode("ascii")
        + b" /Filter /FlateDecode>>\nstream\n"
        + compressed
        + b"\nendstream\nendobj\nstartxref\n0\n%%EOF\n"
    )
    monkeypatch.setattr(cv_importer.fitz, "open", lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("PyMuPDF reached")))
    with pytest.raises(CVImportError, match="compression"):
        cv_importer._preflight_pdf_raw(payload, check_compression=True)


def test_raster_pixel_bomb_rejected_before_pixmap_allocation(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    monkeypatch.setattr(cv_importer, "MAX_PAGE_PIXELS", 1_000_000)

    class FakePage:
        rect = types.SimpleNamespace(width=2_000.0, height=2_000.0)

        def get_pixmap(self, **_kwargs):
            raise AssertionError("pixmap allocation reached")

    with pytest.raises(CVImportError, match="pixel limits"):
        cv_importer._raster_page_png(FakePage(), 1)


def test_isolated_worker_timeout_terminates_tree_without_launch(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    events: list[str] = []

    class FakeGate:
        def set(self):
            events.append("gate")

    class FakeProcess:
        pid = 123
        exitcode = None

        def start(self):
            events.append("start")

        def join(self, _timeout=None):
            events.append("join")

        def is_alive(self):
            return True

    class FakeContext:
        def Event(self):
            return FakeGate()

        def Process(self, **_kwargs):
            return FakeProcess()

    monkeypatch.setattr(cv_importer.multiprocessing, "get_context", lambda _name: FakeContext())
    monkeypatch.setattr(cv_importer, "_windows_job_for_process", lambda _pid, _memory: 7)
    monkeypatch.setattr(cv_importer, "_terminate_worker_tree", lambda _process, job: events.append(f"terminate-{job}"))
    monkeypatch.setattr(cv_importer, "_close_worker_job", lambda job: events.append(f"close-{job}"))
    assert cv_importer._run_isolated_worker(lambda: None, (), timeout_seconds=1, memory_bytes=1024) is False
    assert "terminate-7" in events and events[-1] == "close-7"


def test_windows_job_setup_exception_cleans_spawned_worker_before_gate_release(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    events: list[str] = []

    class FakeGate:
        released = False

        def set(self):
            self.released = True
            events.append("gate")

    class FakeProcess:
        pid = 123
        exitcode = None

        def __init__(self):
            self.alive = True

        def start(self):
            events.append("start")

        def join(self, _timeout=None):
            events.append("join")

        def is_alive(self):
            events.append("alive")
            return self.alive

        def terminate(self):
            events.append("terminate")

        def kill(self):
            events.append("kill")
            self.alive = False

    gate = FakeGate()
    process = FakeProcess()

    class FakeContext:
        def Event(self):
            return gate

        def Process(self, **_kwargs):
            return process

    def failed_job_setup(_pid, _memory):
        events.append("job")
        raise OSError("synthetic job setup failure")

    monkeypatch.setattr(cv_importer.os, "name", "nt")
    monkeypatch.setattr(cv_importer.multiprocessing, "get_context", lambda _name: FakeContext())
    monkeypatch.setattr(cv_importer, "_windows_job_for_process", failed_job_setup)
    monkeypatch.setattr(cv_importer, "_close_worker_job", lambda job: events.append(f"close-{job}"))

    assert cv_importer._run_isolated_worker(lambda: None, (), timeout_seconds=1, memory_bytes=1024) is False
    assert gate.released is False
    assert process.alive is False
    assert events[:2] == ["start", "job"]
    assert "terminate" in events and "kill" in events
    assert events[-1] == "close-None"


@pytest.mark.parametrize(
    ("failure", "expected_closed"),
    [
        ("stdout-open", []),
        ("stderr-open", [101]),
        ("stdout-dup2", [101, 102]),
        ("stderr-dup2", [101, 102]),
    ],
)
def test_spawn_worker_wrapper_fails_closed_when_fd_redirection_fails(
    monkeypatch: pytest.MonkeyPatch, failure: str, expected_closed: list[int]
):
    from Pipeline import cv_importer

    opened: list[int] = []
    closed: list[int] = []
    target_called = False
    original_stdout, original_stderr = sys.stdout, sys.stderr

    def fake_open(path, flags):
        assert path == os.devnull
        assert flags & os.O_WRONLY
        descriptor = 101 + len(opened)
        opened.append(descriptor)
        if failure == "stdout-open" and len(opened) == 1:
            raise OSError("synthetic stdout open failure")
        if failure == "stderr-open" and len(opened) == 2:
            raise OSError("synthetic stderr open failure")
        return descriptor

    def fake_dup2(_source, destination):
        if failure == "stdout-dup2" and destination == 1:
            raise OSError("synthetic stdout redirect failure")
        if failure == "stderr-dup2" and destination == 2:
            raise OSError("synthetic stderr redirect failure")

    def target(_gate):
        nonlocal target_called
        target_called = True

    try:
        with monkeypatch.context() as scoped:
            scoped.setattr(cv_importer.os, "open", fake_open)
            scoped.setattr(cv_importer.os, "dup2", fake_dup2)
            scoped.setattr(cv_importer.os, "close", lambda descriptor: closed.append(descriptor))
            with pytest.raises(SystemExit) as error:
                cv_importer._isolated_worker_entry(target, (), object())
    finally:
        sys.stdout, sys.stderr = original_stdout, original_stderr

    assert error.value.code == 70
    assert target_called is False
    assert closed == expected_closed


def test_spawn_worker_wrapper_suppresses_unexpected_exception_and_cli_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capfd: pytest.CaptureFixture[str]
):
    from Pipeline import cv_importer

    secret = "SYNTHETIC-SECRET-PATH-CONTENT"

    def failed_registration(_root: Path):
        succeeded = cv_importer._run_isolated_worker(
            leaking_worker_target,
            (secret,),
            timeout_seconds=5,
            memory_bytes=cv_importer.MAX_WORKER_MEMORY_BYTES,
        )
        assert succeeded is False
        raise CVImportError("synthetic worker failure")

    monkeypatch.setattr(cv_importer, "_registered_source", failed_registration)
    monkeypatch.setattr(sys, "argv", ["cv_importer.py", "--root", str(tmp_path)])
    assert cv_importer.main() == 2
    captured = capfd.readouterr()
    assert json.loads(captured.out) == {"ok": False, "error": "CV import failed"}
    assert secret not in captured.out + captured.err


def test_pdf_baseline_is_metadata_and_link_free_with_two_x_pngs(tmp_path: Path):
    path = tmp_path / "master.pdf"
    make_pdf(path, link=True, metadata=True)
    descriptor = _extract_layout_to_work_dir(validate_source_cv(path), tmp_path / "work")
    baseline = tmp_path / "work" / "baseline.pdf"
    image = tmp_path / "work" / "page-001.png"
    assert descriptor.page_count == 1
    with fitz.open(baseline) as document:
        assert not document.metadata.get("author")
        assert document[0].get_links() == []
    pixmap = fitz.Pixmap(image)
    assert pixmap.width == 1190
    assert pixmap.height == 1684


@pytest.mark.parametrize(
    "member",
    [
        "../escape.bin",
        "/absolute.bin",
        "word/vbaProject.bin",
        "word/embeddings/object.bin",
        "customUI/customUI.xml",
        "word/activeX/activeX1.bin",
        "word/payload.exe",
        "EncryptedPackage",
    ],
)
def test_docx_import_rejects_traversal_macro_embedded_active_executable_and_encrypted_members(tmp_path: Path, member: str):
    path = tmp_path / "unsafe.docx"
    make_docx(path)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, b"synthetic")
    with pytest.raises(CVImportError):
        validate_source_cv(path)


@pytest.mark.parametrize("instruction", ["DDE cmd", "DDEAUTO cmd", "INCLUDETEXT https://example.test/a", "INCLUDEPICTURE file:///x", "LINK Excel.Sheet"])
def test_docx_import_rejects_dangerous_field_instructions(tmp_path: Path, instruction: str):
    path = tmp_path / "field.docx"
    document = Document()
    paragraph = document.add_paragraph()
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    paragraph._p.append(field)
    document.save(path)
    with pytest.raises(CVImportError, match="field instruction"):
        validate_source_cv(path)


def test_docx_import_rejects_altchunk_dtd_entity_and_zip_bomb(tmp_path: Path):
    altchunk = tmp_path / "altchunk.docx"
    make_docx(altchunk)
    with zipfile.ZipFile(altchunk) as archive:
        xml = archive.read("word/document.xml")
    rewrite_member(altchunk, "word/document.xml", xml.replace(b"</w:body>", b'<w:altChunk r:id="rId9"/></w:body>'))
    with pytest.raises(CVImportError, match="active content"):
        validate_source_cv(altchunk)

    entity = tmp_path / "entity.docx"
    make_docx(entity)
    rewrite_member(entity, "word/document.xml", b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY y SYSTEM "file:///x">]><x>&y;</x>')
    with pytest.raises(CVImportError, match="DTD or entity"):
        validate_source_cv(entity)

    bomb = tmp_path / "bomb.docx"
    make_docx(bomb)
    with zipfile.ZipFile(bomb, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/huge.xml", b"0" * (9 * 1024 * 1024))
    with pytest.raises(CVImportError, match="archive limits"):
        validate_source_cv(bomb)


@pytest.mark.parametrize(
    "payload",
    [
        b"\xff\xfe" + '<!DOCTYPE x [<!ENTITY y "boom">]><x>&y;</x>'.encode("utf-16-le"),
        b"\xfe\xff" + '<!DOCTYPE x [<!ENTITY y "boom">]><x>&y;</x>'.encode("utf-16-be"),
        b"\xff\xfe\x00\x00" + '<!DOCTYPE x [<!ENTITY y "boom">]><x>&y;</x>'.encode("utf-32-le"),
        b"\x00\x00\xfe\xff" + '<!DOCTYPE x [<!ENTITY y "boom">]><x>&y;</x>'.encode("utf-32-be"),
    ],
)
def test_xml_policy_rejects_utf16_utf32_dtd_entity_encodings(payload: bytes):
    from Pipeline import cv_importer

    with pytest.raises(CVImportError, match="encoding"):
        cv_importer._xml_root(payload, "XML")


def test_xml_policy_accepts_benign_utf8_but_rejects_declared_legacy_and_spaced_dtd():
    from Pipeline import cv_importer

    benign = b'<?xml version="1.0" encoding="UTF-8"?><root><item>synthetic</item></root>'
    assert cv_importer._xml_root(benign, "XML").find("item").text == "synthetic"
    with pytest.raises(CVImportError, match="encoding"):
        cv_importer._xml_root(b'<?xml version="1.0" encoding="ISO-8859-1"?><root/>', "XML")
    with pytest.raises(CVImportError, match="DTD or entity"):
        cv_importer._xml_root(b'<?xml version="1.0"?><!  DoCtYpE root [<!  EnTiTy x "y">]><root/>', "XML")


def test_docx_rejects_split_field_code_and_ole_element(tmp_path: Path):
    split = tmp_path / "split-field.docx"
    document = Document()
    paragraph = document.add_paragraph()
    for text in ("INCLUDE", "TEXT https://example.test/a"):
        field = OxmlElement("w:instrText")
        field.text = text
        paragraph._p.append(field)
    document.save(split)
    with pytest.raises(CVImportError, match="field instruction"):
        validate_source_cv(split)

    ole = tmp_path / "ole.docx"
    make_docx(ole)
    with zipfile.ZipFile(ole) as archive:
        xml = archive.read("word/document.xml")
    rewrite_member(ole, "word/document.xml", xml.replace(b"</w:body>", b'<w:object><o:OLEObject xmlns:o="urn:schemas-microsoft-com:office:office"/></w:object></w:body>'))
    with pytest.raises(CVImportError, match="active content"):
        validate_source_cv(ole)


@pytest.mark.parametrize(
    "fragments",
    [
        ("HYPERLINK https://example.test/profile",),
        ("HYPER", "LINK https://example.test/profile"),
        ("DATABASE synthetic",),
        ("RD synthetic",),
        ("UNKNOWNEXTERNAL synthetic",),
    ],
)
def test_docx_rejects_every_field_instruction_including_split_and_unknown(tmp_path: Path, fragments: tuple[str, ...]):
    path = tmp_path / "field.docx"
    document = Document()
    paragraph = document.add_paragraph()
    for text in fragments:
        field = OxmlElement("w:instrText")
        field.text = text
        paragraph._p.append(field)
    document.save(path)
    with pytest.raises(CVImportError, match="field instruction"):
        validate_source_cv(path)


def test_docx_rejects_xml_encoded_field_instruction(tmp_path: Path):
    path = tmp_path / "encoded-field.docx"
    document = Document()
    document.add_paragraph("Synthetic")
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    rewrite_member(path, "word/document.xml", xml.replace(b"</w:body>", b"<w:p><w:r><w:instrText>HYPER&#76;INK https://example.test</w:instrText></w:r></w:p></w:body>"))
    with pytest.raises(CVImportError, match="field instruction"):
        validate_source_cv(path)


def test_docx_accepts_passive_http_hyperlink_but_rejects_other_external_relationships(tmp_path: Path):
    passive = tmp_path / "passive.docx"
    make_docx(passive)
    add_external_relationship(
        passive,
        relation_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        target="https://example.test/profile",
    )
    assert validate_source_cv(passive).kind == "docx"

    active = tmp_path / "external-template.docx"
    make_docx(active)
    add_external_relationship(
        active,
        relation_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
        target="https://example.test/template.dotm",
    )
    with pytest.raises(CVImportError, match="external relationship"):
        validate_source_cv(active)

    external_file = tmp_path / "external-file-image.docx"
    make_docx(external_file)
    add_external_relationship(
        external_file,
        relation_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        target="file:///synthetic/private.png",
    )
    with pytest.raises(CVImportError, match="external relationship"):
        validate_source_cv(external_file)


def test_docx_rejects_case_colliding_archive_members(tmp_path: Path):
    path = tmp_path / "case-collision.docx"
    make_docx(path)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("WORD/DOCUMENT.XML", b"synthetic")
    with pytest.raises(CVImportError, match="duplicate members"):
        validate_source_cv(path)


@pytest.mark.parametrize(
    ("replacement", "message"),
    [
        (b'<svg xmlns="http://www.w3.org/2000/svg"><script>1</script></svg>', "media"),
        (b"<html>synthetic</html>", "media"),
        (b"<?xml version='1.0'?><x/>", "media"),
        (png_with_dimensions(100_000, 100_000), "dimensions"),
    ],
)
def test_docx_rejects_active_or_oversized_bytes_disguised_as_png(tmp_path: Path, replacement: bytes, message: str):
    path = tmp_path / "bad-media.docx"
    make_docx(path)
    with zipfile.ZipFile(path) as archive:
        media_name = next(name for name in archive.namelist() if name.casefold().endswith(".png") and name.casefold().startswith("word/media/"))
    rewrite_member(path, media_name, replacement)
    with pytest.raises(CVImportError, match=message):
        validate_source_cv(path)


def test_docx_rejects_media_content_type_mismatch_and_unknown_internal_relationship(tmp_path: Path):
    mismatch = tmp_path / "mismatch.docx"
    make_docx(mismatch)
    with zipfile.ZipFile(mismatch) as archive:
        content_types = archive.read("[Content_Types].xml")
    rewrite_member(mismatch, "[Content_Types].xml", content_types.replace(b'Extension="png" ContentType="image/png"', b'Extension="png" ContentType="image/jpeg"'))
    with pytest.raises(CVImportError, match="content type"):
        validate_source_cv(mismatch)

    unknown = tmp_path / "unknown-relationship.docx"
    make_docx(unknown)
    add_internal_relationship(
        unknown,
        relation_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
        target="../customXml/item1.xml",
    )
    with pytest.raises(CVImportError, match="relationship type"):
        validate_source_cv(unknown)


def test_docx_rejects_unallowlisted_content_type_and_lookalike_hyperlink_relationship(tmp_path: Path):
    unknown_type = tmp_path / "unknown-content-type.docx"
    make_docx(unknown_type)
    with zipfile.ZipFile(unknown_type) as archive:
        content_types = archive.read("[Content_Types].xml")
    rewrite_member(
        unknown_type,
        "[Content_Types].xml",
        content_types.replace(
            b"</Types>",
            b'<Default Extension="synthetic" ContentType="text/html"/></Types>',
        ),
    )
    with zipfile.ZipFile(unknown_type, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/unreferenced.synthetic", b"<html>active</html>")
    with pytest.raises(CVImportError, match="content type"):
        validate_source_cv(unknown_type)

    lookalike = tmp_path / "lookalike-hyperlink.docx"
    make_docx(lookalike)
    add_external_relationship(
        lookalike,
        relation_type="https://attacker.invalid/relationships/hyperlink",
        target="https://example.test/profile",
    )
    with pytest.raises(CVImportError, match="external relationship"):
        validate_source_cv(lookalike)


def test_docx_metadata_limits_are_checked_before_any_member_read(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    content_types = zipfile.ZipInfo("[Content_Types].xml")
    content_types.file_size = 100
    content_types.compress_size = 80
    content_types.compress_type = zipfile.ZIP_DEFLATED
    bomb = zipfile.ZipInfo("word/bomb.xml")
    bomb.file_size = 5 * 1024 * 1024
    bomb.compress_size = 1
    bomb.compress_type = zipfile.ZIP_DEFLATED
    reads: list[str] = []

    class MetadataOnlyArchive:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def infolist(self):
            return [content_types, bomb]

        def read(self, member):
            reads.append(str(member))
            raise AssertionError("archive member decompression reached")

    monkeypatch.setattr(cv_importer.zipfile, "ZipFile", lambda *_args, **_kwargs: MetadataOnlyArchive())
    with pytest.raises(CVImportError, match="archive limits"):
        cv_importer._validate_docx_payload(b"PK\x03\x04synthetic")
    assert reads == []


def test_docx_rejects_untyped_part_and_image_outside_word_media(tmp_path: Path):
    untyped = tmp_path / "untyped.docx"
    make_docx(untyped)
    with zipfile.ZipFile(untyped, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/untyped", b"synthetic")
    with pytest.raises(CVImportError, match="content type"):
        validate_source_cv(untyped)

    outside_image = tmp_path / "outside-image.docx"
    make_docx(outside_image)
    with zipfile.ZipFile(outside_image, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("customXml/disguised.png", b"<svg><script>1</script></svg>")
    with pytest.raises(CVImportError, match="media"):
        validate_source_cv(outside_image)


def test_docx_media_decoded_byte_limit_precedes_image_decode(monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    class OversizedDecodedImage:
        size = (5_000, 4_000)
        format = "PNG"

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def getbands(self):
            return ("R", "G", "B", "A")

        def verify(self):
            raise AssertionError("image decode reached")

    monkeypatch.setattr(cv_importer.Image, "open", lambda *_args, **_kwargs: OversizedDecodedImage())
    with pytest.raises(CVImportError, match="decoded size"):
        cv_importer._validate_media("customXml/image.png", PNG_1X1, "image/png")


def test_docx_image_relationship_must_target_validated_image_part(tmp_path: Path):
    path = tmp_path / "wrong-image-target.docx"
    make_docx(path)
    add_internal_relationship(
        path,
        relation_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        target="../customXml/item1.xml",
    )
    with pytest.raises(CVImportError, match="relationship target type"):
        validate_source_cv(path)


def test_docx_requires_exact_canonical_content_types_member_name(tmp_path: Path):
    path = tmp_path / "wrong-case-content-types.docx"
    make_docx(path)
    rename_member(path, "[Content_Types].xml", "[content_types].xml")
    with pytest.raises(CVImportError, match="content type"):
        validate_source_cv(path)


def test_docx_rejects_macro_content_type_variant(tmp_path: Path):
    path = tmp_path / "macro-type.docx"
    make_docx(path)
    with zipfile.ZipFile(path) as archive:
        content_types = archive.read("[Content_Types].xml")
    rewrite_member(
        path,
        "[Content_Types].xml",
        content_types.replace(
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            b"application/vnd.ms-word.document.macroEnabled.main+xml",
        ),
    )
    with pytest.raises(CVImportError):
        validate_source_cv(path)


def test_docx_no_renderer_uses_exact_pdf_retry_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    path = tmp_path / "master.docx"
    make_docx(path)
    monkeypatch.setattr(cv_importer, "_find_libreoffice", lambda: None)
    monkeypatch.setattr(cv_importer, "_render_docx_with_word", lambda *_args: None)
    with pytest.raises(
        CVImportError,
        match="^DOCX visual import requires an available isolated LibreOffice renderer; export the master CV as PDF and retry$",
    ):
        _extract_layout_to_work_dir(validate_source_cv(path), tmp_path / "work")


def test_docx_converted_active_pdf_is_rejected_before_layout_extraction(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    path = tmp_path / "master.docx"
    make_docx(path)

    def active_renderer(_payload: bytes, work_dir: Path) -> bytes:
        output = work_dir / "converted.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Synthetic")
        document.xref_set_key(page.xref, "AA", "<</O<</S/JavaScript/JS(app.alert('x'))>>>>")
        document.save(output, deflate=True, garbage=4, use_objstms=1)
        document.close()
        return output.read_bytes()

    monkeypatch.setattr(cv_importer, "_render_docx_to_pdf", active_renderer)
    monkeypatch.setattr(
        cv_importer,
        "_extract_pdf_isolated",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("layout extraction reached")),
    )
    with pytest.raises(CVImportError, match="active content"):
        _extract_layout_to_work_dir(validate_source_cv(path), tmp_path / "work")


def test_libreoffice_conversion_delegates_to_unique_bounded_isolated_worker(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    work = tmp_path / "private-work"
    work.mkdir()
    source = work / "source.docx"
    source.write_bytes(b"synthetic")
    executable = work / "soffice.exe"
    executable.write_bytes(b"synthetic")
    invocation_paths: list[Path] = []

    def fake_runner(target, arguments, *, timeout_seconds, memory_bytes):
        assert target is cv_importer._libreoffice_worker
        assert timeout_seconds == cv_importer.MAX_RENDER_SECONDS
        assert 0 < memory_bytes <= 512 * 1024 * 1024
        assert Path(arguments[0]) == executable
        private_source = Path(arguments[1])
        output = Path(arguments[2])
        isolation = Path(arguments[3])
        assert private_source.parent == isolation and output.parent == isolation
        assert isolation.is_relative_to(work) and isolation.name.startswith("libreoffice-")
        invocation_paths.append(isolation)
        make_pdf(output)
        return True

    monkeypatch.setattr(cv_importer, "_run_isolated_worker", fake_runner, raising=False)
    monkeypatch.setattr(
        cv_importer.subprocess,
        "run",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("direct subprocess launch reached")),
    )
    payload = cv_importer._render_docx_with_libreoffice(executable, source, work)
    assert payload is not None and payload.startswith(b"%PDF-")
    assert invocation_paths and not invocation_paths[0].exists()


def test_libreoffice_worker_uses_devnull_and_minimal_environment(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    isolation = tmp_path / "private-libreoffice"
    isolation.mkdir()
    executable = isolation / "soffice.exe"
    executable.write_bytes(b"synthetic")
    source = isolation / "input.docx"
    source.write_bytes(b"synthetic")
    output = isolation / "input.pdf"
    monkeypatch.setenv("SYNTHETIC_CREDENTIAL_TOKEN", "must-not-be-inherited")
    monkeypatch.setattr(cv_importer, "_apply_posix_worker_limits", lambda *_args: None)

    def fake_run(command, **options):
        assert command[0] == str(executable)
        assert options["stdin"] is subprocess.DEVNULL
        assert options["stdout"] is subprocess.DEVNULL
        assert options["stderr"] is subprocess.DEVNULL
        assert options["shell"] is False
        assert "SYNTHETIC_CREDENTIAL_TOKEN" not in options["env"]
        assert set(options["env"]).issubset({"PATH", "SystemRoot", "TEMP", "TMP", "HOME", "USERPROFILE", "APPDATA", "LOCALAPPDATA", "SAL_DISABLE_OPENCL", "SAL_DISABLEGL"})
        make_pdf(output)
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(cv_importer.subprocess, "run", fake_run)
    cv_importer._libreoffice_worker(
        str(executable), str(source), str(output), str(isolation), types.SimpleNamespace(wait=lambda: True)
    )
    assert output.is_file()


def test_word_fallback_is_disabled_without_reliable_os_network_boundary(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    work = tmp_path / "private-work"
    work.mkdir()
    source = work / "source.docx"
    source.write_bytes(b"synthetic")
    output = work / "word.pdf"
    monkeypatch.setattr(
        cv_importer,
        "_run_isolated_worker",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("Word worker reached")),
    )
    assert cv_importer._render_docx_with_word(source, output, work) is None
    assert not output.exists()


def test_docx_fake_renderer_extracts_paragraph_styles_tables_sections_headers_footers_images_and_order(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    path = tmp_path / "master.docx"
    make_docx(path)
    monkeypatch.setattr(cv_importer, "_render_docx_to_pdf", fake_docx_renderer)
    descriptor = _extract_layout_to_work_dir(validate_source_cv(path), tmp_path / "work")
    types = {block.block_type for block in descriptor.blocks}
    texts = {block.text for block in descriptor.blocks}
    assert {"paragraph", "table", "section", "header", "footer"}.issubset(types)
    assert {"Synthetic role | 2030", "Synthetic header", "Synthetic footer"}.issubset(texts)
    assert any("Synthetic Heading" in block.style_names for block in descriptor.blocks)
    assert descriptor.pages[0].margin_top_points == pytest.approx(36)
    assert descriptor.pages[0].margin_right_points == pytest.approx(57.6)
    assert descriptor.image_count >= 1
    assert "EXPERIENCE" in descriptor.section_order
    assert (tmp_path / "work" / "baseline.pdf").is_file()
    assert (tmp_path / "work" / "page-001.png").is_file()
    assert not (tmp_path / "work" / "source.docx").exists()


def test_work_directory_rejects_link_escape_and_leaves_outside_empty(tmp_path: Path):
    source = tmp_path / "master.pdf"
    make_pdf(source)
    outside = tmp_path / "outside"
    outside.mkdir()
    linked = tmp_path / "linked-work"
    try:
        linked.symlink_to(outside, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"symlink unavailable: {error}")
    with pytest.raises(CVImportError, match="work directory"):
        _extract_layout_to_work_dir(validate_source_cv(source), linked)
    assert list(outside.iterdir()) == []


def test_docx_renderer_rejects_linked_unique_workspace_before_worker_launch(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    source_path = tmp_path / "master.docx"
    make_docx(source_path)
    work = tmp_path / "work"
    work.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    try:
        (work / "libreoffice-fixed").symlink_to(outside, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"symlink unavailable: {error}")
    executable = work / "soffice.exe"
    executable.write_bytes(b"synthetic")
    monkeypatch.setattr(cv_importer, "_find_libreoffice", lambda: executable)
    monkeypatch.setattr(cv_importer.uuid, "uuid4", lambda: types.SimpleNamespace(hex="fixed"))
    monkeypatch.setattr(
        cv_importer,
        "_run_isolated_worker",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("renderer worker reached")),
    )
    with pytest.raises(CVImportError):
        _extract_layout_to_work_dir(validate_source_cv(source_path), work)
    assert list(outside.iterdir()) == []


def registered_descriptor(project_root: Path):
    initialize_project(project_root)
    source_path = project_root / "source.pdf"
    make_pdf(source_path)
    registration = register_master_cv(project_root, source_path)
    source = validate_source_cv(project_root / registration["path"])
    return extract_layout(source, project_root)


def test_public_extract_writes_only_registered_project_baseline(project_root: Path):
    initialize_project(project_root)
    source_path = project_root / "source.pdf"
    make_pdf(source_path)
    registration = register_master_cv(project_root, source_path)
    source = validate_source_cv(project_root / registration["path"])

    descriptor = extract_layout(source, project_root)

    baseline = project_root / "private" / "cv-creator" / "baseline"
    assert descriptor.source_sha256 == registration["sha256"]
    assert (baseline / "baseline.pdf").is_file()
    assert (baseline / "page-001.png").is_file()
    assert not (project_root / "baseline.pdf").exists()


def test_public_extract_rejects_arbitrary_unregistered_root(tmp_path: Path):
    source_path = tmp_path / "source.pdf"
    make_pdf(source_path)
    outside = tmp_path / "chosen-output"

    with pytest.raises(CVImportError, match="registered master CV"):
        extract_layout(validate_source_cv(source_path), outside)

    assert not outside.exists()


def test_layout_descriptor_writes_only_exact_private_path_and_contains_no_absolute_source_path(project_root: Path):
    descriptor = registered_descriptor(project_root)
    output = write_layout_descriptor(project_root, descriptor)
    assert output == project_root / "private" / "cv-creator" / "layout.json"
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["source_path"] == "private/master-cv.pdf"
    serialized = output.read_text(encoding="utf-8")
    assert str(project_root) not in serialized
    assert "source.pdf" not in serialized
    assert not list(output.parent.glob("*.tmp-*"))


@pytest.mark.parametrize("mutation", ["hash", "page_total", "page_number", "bbox", "image_total", "kind"])
def test_layout_descriptor_rejects_drifted_or_internally_inconsistent_values(project_root: Path, mutation: str):
    descriptor = registered_descriptor(project_root)
    if mutation == "hash":
        descriptor = replace(descriptor, source_sha256="0" * 64)
    elif mutation == "page_total":
        descriptor = replace(descriptor, page_count=2)
    elif mutation == "page_number":
        descriptor = replace(descriptor, pages=(replace(descriptor.pages[0], page_number=2),))
    elif mutation == "bbox":
        descriptor = replace(descriptor, blocks=(replace(descriptor.blocks[0], bbox=(-1.0, 0.0, 1.0, 1.0)),) + descriptor.blocks[1:])
    elif mutation == "image_total":
        descriptor = replace(descriptor, image_count=999)
    else:
        descriptor = replace(descriptor, source_kind="docx")
    with pytest.raises(CVImportError, match="layout descriptor"):
        write_layout_descriptor(project_root, descriptor)
    assert not (project_root / "private" / "cv-creator" / "layout.json").exists()


def test_layout_descriptor_rejects_geometry_drift_on_later_page(project_root: Path):
    initialize_project(project_root)
    source_path = project_root / "source.pdf"
    make_pdf(source_path, pages=2)
    registration = register_master_cv(project_root, source_path)
    source = validate_source_cv(project_root / registration["path"])
    descriptor = extract_layout(source, project_root)
    changed_page = replace(descriptor.pages[1], width_points=descriptor.pages[1].width_points - 1)
    changed = replace(descriptor, pages=(descriptor.pages[0], changed_page))
    with pytest.raises(CVImportError, match="layout descriptor"):
        write_layout_descriptor(project_root, changed)


def test_forged_docx_page_total_cannot_reuse_same_run_attestation(project_root: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    initialize_project(project_root)
    source_path = project_root / "source.docx"
    make_docx(source_path)
    registration = register_master_cv(project_root, source_path)
    monkeypatch.setattr(cv_importer, "_render_docx_to_pdf", fake_docx_renderer)
    descriptor = extract_layout(validate_source_cv(project_root / registration["path"]), project_root)
    pages = tuple(replace(descriptor.pages[0], page_number=number) for number in range(1, 5))
    forged = replace(descriptor, page_count=4, pages=pages)

    with pytest.raises(CVImportError, match="layout descriptor"):
        write_layout_descriptor(project_root, forged)

    assert not (project_root / "private" / "cv-creator" / "layout.json").exists()


def test_new_extraction_invalidates_stale_descriptor_attestation(project_root: Path):
    stale = registered_descriptor(project_root)
    source = validate_source_cv(project_root / "private" / "master-cv.pdf")
    current = extract_layout(source, project_root)

    with pytest.raises(CVImportError, match="attestation"):
        write_layout_descriptor(project_root, stale)

    assert write_layout_descriptor(project_root, current).is_file()


def test_descriptor_token_invalidation_between_precheck_and_project_lock_fails_closed(
    project_root: Path, monkeypatch: pytest.MonkeyPatch
):
    from Pipeline import cv_importer

    descriptor = registered_descriptor(project_root)
    real_lock = cv_importer.exclusive_project_lock

    @contextmanager
    def invalidating_lock(root: Path, name: str):
        with real_lock(root, name=name):
            with cv_importer._ATTESTATION_LOCK:
                cv_importer._ATTESTATIONS.clear()
            yield

    monkeypatch.setattr(cv_importer, "exclusive_project_lock", invalidating_lock)
    with pytest.raises(CVImportError, match="attestation"):
        write_layout_descriptor(project_root, descriptor)
    assert not (project_root / "private" / "cv-creator" / "layout.json").exists()


def test_extraction_publication_uses_same_project_import_lock(project_root: Path, monkeypatch: pytest.MonkeyPatch):
    from Pipeline import cv_importer

    initialize_project(project_root)
    source_path = project_root / "source.pdf"
    make_pdf(source_path)
    registration = register_master_cv(project_root, source_path)
    source = validate_source_cv(project_root / registration["path"])
    entered: list[str] = []

    @contextmanager
    def recording_lock(_root: Path, name: str):
        entered.append(name)
        yield

    monkeypatch.setattr(cv_importer, "exclusive_project_lock", recording_lock)
    extract_layout(source, project_root)
    assert entered == [".cv-import.lock"]


def test_successful_descriptor_attestation_cannot_be_replayed(project_root: Path):
    descriptor = registered_descriptor(project_root)
    output = write_layout_descriptor(project_root, descriptor)
    committed = output.read_bytes()

    with pytest.raises(CVImportError, match="attestation"):
        write_layout_descriptor(project_root, descriptor)

    assert output.read_bytes() == committed


@pytest.mark.parametrize("axis", ["horizontal", "vertical"])
def test_descriptor_rejects_margin_sums_beyond_page(axis: str):
    from Pipeline import cv_importer

    page = cv_importer.PageGeometry(1, 100.0, 200.0)
    if axis == "horizontal":
        page = replace(page, margin_left_points=60.0, margin_right_points=60.0)
    else:
        page = replace(page, margin_top_points=110.0, margin_bottom_points=110.0)
    descriptor = cv_importer.LayoutDescriptor("a" * 64, "pdf", 1, (page,), (), (), 0)

    with pytest.raises(CVImportError, match="margins"):
        cv_importer._validate_descriptor_shape(descriptor)


@pytest.mark.parametrize("mutation", ["runtime", "source"])
def test_layout_commit_rolls_back_when_registered_identity_changes(
    project_root: Path, monkeypatch: pytest.MonkeyPatch, mutation: str
):
    from Pipeline import cv_importer

    descriptor = registered_descriptor(project_root)
    original_write = cv_importer.atomic_write_json
    triggered = False

    def racing_write(path, value, project_root=None):
        nonlocal triggered
        original_write(path, value, project_root=project_root)
        if Path(path).name != "layout.json" or triggered:
            return
        triggered = True
        if mutation == "runtime":
            runtime_path = project_root / "private" / "runtime.json"
            runtime = json.loads(runtime_path.read_text(encoding="utf-8"))
            runtime["master_cv_sha256"] = "0" * 64
            original_write(runtime_path, runtime, project_root=project_root)
        else:
            replacement = project_root / "replacement.pdf"
            make_pdf(replacement, pages=2)
            os.replace(replacement, project_root / "private" / "master-cv.pdf")

    monkeypatch.setattr(cv_importer, "atomic_write_json", racing_write)
    with pytest.raises(CVImportError, match="changed during layout commit"):
        write_layout_descriptor(project_root, descriptor)
    assert not (project_root / "private" / "cv-creator" / "layout.json").exists()


@pytest.mark.parametrize("link_kind", ["symlink", "hardlink"])
def test_layout_descriptor_rejects_link_like_destination(project_root: Path, tmp_path: Path, link_kind: str):
    descriptor = registered_descriptor(project_root)
    destination = project_root / "private" / "cv-creator" / "layout.json"
    outside = tmp_path / "outside.json"
    outside.write_text("outside", encoding="utf-8")
    try:
        if link_kind == "symlink":
            destination.symlink_to(outside)
        else:
            os.link(outside, destination)
    except OSError as error:
        pytest.skip(f"link unavailable: {error}")
    with pytest.raises(CVImportError, match="layout destination"):
        write_layout_descriptor(project_root, descriptor)
    assert outside.read_text(encoding="utf-8") == "outside"


def test_cli_uses_registered_runtime_cv_and_returns_field_safe_status(project_root: Path):
    initialize_project(project_root)
    source = project_root / "SYNTHETIC-PRIVATE-FILENAME.pdf"
    make_pdf(source)
    register_master_cv(project_root, source)
    result = subprocess.run(
        [sys.executable, "Pipeline/cv_importer.py", "--root", "."],
        cwd=project_root,
        text=True,
        capture_output=True,
        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        timeout=30,
        check=False,
    )
    assert result.returncode == 0
    status = json.loads(result.stdout)
    assert status == {"ok": True, "kind": "pdf", "page_count": 1}
    combined = result.stdout + result.stderr
    assert "SYNTHETIC" not in combined
    assert str(project_root) not in combined
    assert (project_root / "private" / "cv-creator" / "layout.json").is_file()
